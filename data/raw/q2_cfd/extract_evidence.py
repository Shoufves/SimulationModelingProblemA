"""Extract award terms without promoting them to national or actual generation."""
from collections import Counter
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import openpyxl

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
manifest=json.loads((HERE/'source_manifest.json').read_text('utf-8'))
for entry in manifest:
    if hashlib.sha256((ROOT/entry['path']).read_bytes()).hexdigest()!=entry['sha256']:
        raise ValueError('Original changed: '+entry['path'])
workbook=openpyxl.load_workbook(HERE/'shanghai_2025_awards.xlsx',data_only=True,read_only=True)
sheet=workbook['Sheet0']
headers=[cell.value for cell in next(sheet.iter_rows())]
if headers[7:]!=['机制电量（兆瓦时)','机制电价（元/千瓦时)','机制电价执行期限(年)']:
    raise ValueError('Unexpected units/header')
rows=[];ids=set();types=Counter();total=Decimal('0')
for index,values in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
    if len(values)!=10 or not values[1] or values[1] in ids:
        raise ValueError('Missing/duplicate project identifier')
    ids.add(values[1]);types[values[2]]+=1
    volume,price,years=(Decimal(str(values[j])) for j in (7,8,9))
    if not all(v.is_finite() and v>0 for v in (volume,price,years)):
        raise ValueError('Invalid award value')
    total+=volume
    rows.append(dict(sheet=sheet.title,row=index,project_type=values[2],
        promised_commissioning_date=str(values[5]),mechanism_volume_mwh=str(volume),
        mechanism_price_nominal_vat_inclusive_cny_per_kwh=str(price),term_years=str(years),
        cells=dict(volume=f'H{index}',price=f'I{index}',term=f'J{index}',promised_date=f'F{index}')))
summary=dict(scope='Shanghai 2025 selected incremental projects only; not national weights or actual generation',
    notification_date='2025-12-03',earliest_settlement_for_already_commissioned='2026-01-01',
    future_projects_settlement='first day of month following promised commissioning, subject to original rules',
    project_count=len(rows),types=dict(types),total_mechanism_volume_mwh=str(total),
    prices_cny_per_kwh=sorted({r['mechanism_price_nominal_vat_inclusive_cny_per_kwh'] for r in rows}),
    terms_years=sorted({r['term_years'] for r in rows}),
    national_model_activated=False,observed_receipts=False,observed_commissioning=False,
    source_sha256={r['path']:r['sha256'] for r in manifest},rows=rows)
if summary['project_count']!=2039 or summary['prices_cny_per_kwh']!=['0.4155']:
    raise ValueError('Workbook disagrees with notification')
(HERE/'awards_evidence.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n','utf-8')
print(json.dumps({k:v for k,v in summary.items() if k not in {'rows','source_sha256'}},ensure_ascii=False))
