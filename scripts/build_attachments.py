"""Read contest attachments, preserve row provenance, and build limited-use asset profiles."""
import json
import math
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from workflow_status import sha256, check_hash

ROOT = Path(__file__).resolve().parents[1]
SHEETS = [
    ("附件1.xlsx", "火力发电厂", "ATT1-THERMAL", "thermal"),
    ("附件2.xlsx", "水电站", "ATT2-HYDRO", "hydro"),
    ("附件2.xlsx", "抽水蓄能电站", "ATT2-PUMP", "pumped_storage"),
    ("附件3.xlsx", "风电场", "ATT3-WIND", "wind"),
    ("附件4.xlsx", "核电站", "ATT4-NUCLEAR", "nuclear"),
]
PROVINCES = ["北京", "天津", "河北", "山西", "内蒙古", "辽宁", "吉林", "黑龙江",
             "上海", "江苏", "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北",
             "湖南", "广东", "广西", "海南", "重庆", "四川", "贵州", "云南", "西藏",
             "陕西", "甘肃", "青海", "宁夏", "新疆"]
SPECIAL = {"内蒙古": "内蒙古自治区", "广西": "广西壮族自治区", "西藏": "西藏自治区",
           "宁夏": "宁夏回族自治区", "新疆": "新疆维吾尔自治区"}
MUNICIPAL = {"北京", "天津", "上海", "重庆"}


def clean(value):
    return value.strip() if isinstance(value, str) else value


def numeric(value):
    return type(value) in (int, float) and math.isfinite(value)


def explicit_province(value):
    """Normalize explicit province prefixes only; never infer from a county name."""
    if not isinstance(value, str):
        return None
    for prefix in sorted(PROVINCES, key=len, reverse=True):
        if value.strip().startswith(prefix):
            return SPECIAL.get(prefix, prefix + ("市" if prefix in MUNICIPAL else "省"))
    return None


def formula_kw(capacity, formula, row):
    """Only interpret the exact supplied multiplication pattern. Never eval formulas."""
    if numeric(capacity) and formula == f"=B{row}*10^3":
        return capacity * 1000
    return None


def tag_duplicates(records):
    groups = defaultdict(list)
    for record in records:
        groups[(record["source_file"], record["sheet"], record["name"])].append(record)
    duplicates = []
    for (_, _, name), items in groups.items():
        if len(items) < 2:
            continue
        signatures = {(r["original_capacity"], r["location"]) for r in items}
        kind = "same_name_same_capacity_location" if len(signatures) == 1 else "same_name_conflicting_fields"
        group_id = items[0]["record_id"] + "-DUP"
        for item in items:
            item["duplicate_group"] = group_id
            item["issues"].append(kind)
        duplicates.append({
            "group_id": group_id, "name": name, "kind": kind,
            "records": [{"record_id": r["record_id"], "row": r["row"],
                         "capacity_raw": r["original_capacity"], "location": r["location"]}
                        for r in items],
            "decision": "retain_all_pending_project_or_unit_identity_review",
        })
    return duplicates


def apply_reviews(records, root):
    config = json.loads((root / "data/attachment_reviews.json").read_text(encoding="utf-8"))
    by_id = {r["record_id"]: r for r in records}
    registry = {s["source_id"]: s for s in json.loads(
        (root / "data/sources.json").read_text(encoding="utf-8"))["sources"]}
    for review in config["reviews"]:
        record = by_id[review["record_id"]]
        if record["name"] != review["expected_name"]:
            raise ValueError("Attachment identity changed: " + record["record_id"])
        check_hash(root, review["evidence_path"], review["evidence_sha256"])
        if review["source_id"] not in registry:
            raise ValueError("Unregistered attachment evidence")
        if "verified_province" in review:
            if record["original_province"] != review["original_province"]:
                raise ValueError("Original province changed: " + record["record_id"])
            record["province"] = review["verified_province"]
            record["province_basis"] = "official_document_field_review"
            record["issues"].append("province_corrected_with_official_evidence")
        record["review_source_ids"].append(review["source_id"])
        record["processing_reason"].append(review["reason"])
        if review.get("coordinate_action") == "quarantine":
            record["coordinates"] = None
            record["coordinate_status"] = "quarantined_pending_location_verification"
            record["issues"].append("coordinates_quarantined")
        if review.get("capacity_reference"):
            record["official_capacity_reference"] = review["capacity_reference"]
        record["verification_status"] = "partially_field_reviewed"
    return config


def extract(root=ROOT):
    originals = json.loads((root / "workflow/provided_inputs.json").read_text(encoding="utf-8"))
    hashes = {r["path"]: r["sha256"] for r in originals["inputs"]}
    records, sheet_info, auxiliary = [], [], []
    for filename, sheet_name, prefix, technology in SHEETS:
        check_hash(root, filename, hashes[filename])
        formulas = openpyxl.load_workbook(root / filename, data_only=False, read_only=True)
        values = openpyxl.load_workbook(root / filename, data_only=True, read_only=True)
        try:
            sheet, cached = formulas[sheet_name], values[sheet_name]
            original_rows = list(sheet.iter_rows(values_only=True))
            cached_rows = list(cached.iter_rows(values_only=True))
            headers = original_rows[0]
            named = 0
            blank_rows = []
            for row_no, cells in enumerate(original_rows[1:], 2):
                cached_cells = cached_rows[row_no - 1]
                if clean(cells[0]) in (None, ""):
                    blank_rows.append(row_no)
                    for col, value in enumerate(cells, 1):
                        if value is not None:
                            auxiliary.append({"source_file": filename, "sheet": sheet_name,
                                              "row": row_no, "column": col, "value": value,
                                              "cached_value": cached_cells[col - 1],
                                              "use": "excluded_unnamed_row"})
                    continue
                named += 1
                raw_cells = {}
                for col, value in enumerate(cells, 1):
                    if value is not None:
                        coordinate = openpyxl.utils.get_column_letter(col) + str(row_no)
                        raw_cells[coordinate] = {
                            "original": value, "cached": cached_cells[col - 1]}
                province_raw = cells[3] if technology in {"pumped_storage", "wind"} else None
                province = explicit_province(province_raw) if province_raw else explicit_province(cells[2])
                lon = cells[4] if technology in {"pumped_storage", "wind"} else None
                lat = cells[5] if technology in {"pumped_storage", "wind"} else None
                has_coords = lon is not None or lat is not None
                coords_valid = numeric(lon) and numeric(lat) and -180 <= lon <= 180 and -90 <= lat <= 90
                computed_kw = formula_kw(cells[1], cells[3], row_no) if technology == "thermal" else None
                cached_kw = cached_cells[3] if technology == "thermal" else None
                kw_label = headers[3] if technology == "thermal" else None
                issues = ["reference_year_missing", "operating_status_missing"]
                if province is None:
                    issues.append("province_unresolved")
                if technology == "thermal":
                    issues.append("thermal_fuel_unresolved")
                    if computed_kw is None:
                        issues.append("unrecognized_kw_formula")
                    elif not numeric(cached_kw) or not math.isclose(computed_kw, cached_kw, abs_tol=0.01):
                        issues.append("kw_cached_value_mismatch")
                else:
                    issues.append("capacity_unit_not_declared")
                if has_coords:
                    issues.append("coordinate_reference_system_unknown")
                    issues.append("coordinates_not_independently_verified")
                    if not coords_valid:
                        issues.append("coordinates_out_of_range_or_missing")
                if technology == "pumped_storage":
                    issues.append("storage_energy_capacity_unknown")
                record = {
                    "record_id": f"{prefix}-R{row_no:04d}",
                    "source_file": filename, "sheet": sheet_name, "row": row_no,
                    "original_name": cells[0], "name": clean(cells[0]), "technology": technology,
                    "original_capacity": cells[1], "original_unit": None,
                    "declared_kw_header": kw_label, "kw_formula": cells[3] if technology == "thermal" else None,
                    "kw_cached_value": cached_kw, "kw_recomputed": computed_kw,
                    "capacity_mw_from_declared_kw": computed_kw / 1000 if computed_kw is not None and kw_label == "KW" else None,
                    "capacity_interpretation": "source_formula_and_KW_header_only" if computed_kw is not None else "unit_unresolved",
                    "location": clean(cells[2]), "original_province": province_raw,
                    "province": province,
                    "province_basis": ("provided_province_column_unverified" if province_raw else
                                       "explicit_location_prefix_unverified" if province else "unknown"),
                    "original_coordinates": {"longitude": lon, "latitude": lat} if has_coords else None,
                    "coordinates": {"longitude": lon, "latitude": lat} if coords_valid else None,
                    "coordinate_status": "provided_unverified" if coords_valid else "missing_or_invalid",
                    "coordinate_reference_system": None, "year": None, "operating_status": None,
                    "official_capacity_reference": None, "duplicate_group": None,
                    "verification_status": "source_transcribed_pending_external_checks",
                    "review_source_ids": [], "issues": issues,
                    "processing_reason": ["保留原行、原值与公式；无时间和状态依据时不补造。",
                                          "不得直接进入全国装机汇总或历史校准。"],
                    "national_calibration_eligible": False, "raw_cells": raw_cells,
                }
                records.append(record)
                used_columns = {1, 2, 3, 4} if technology == "thermal" else (
                    {1, 2, 3, 4, 5, 6} if technology in {"pumped_storage", "wind"} else {1, 2, 3})
                for col, value in enumerate(cells, 1):
                    if col not in used_columns and value is not None:
                        auxiliary.append({"source_file": filename, "sheet": sheet_name, "row": row_no,
                                          "column": col, "value": value, "cached_value": cached_cells[col - 1],
                                          "use": "excluded_auxiliary_cell"})
            sheet_info.append({"source_file": filename, "sheet": sheet_name, "technology": technology,
                               "max_row": sheet.max_row, "named_records": named,
                               "ignored_blank_name_rows": blank_rows, "headers": list(headers),
                               "source_sha256": hashes[filename].lower()})
        finally:
            formulas.close()
            values.close()
    duplicates = tag_duplicates(records)
    review_config = apply_reviews(records, root)
    for filename in {x[0] for x in SHEETS}:
        check_hash(root, filename, hashes[filename])
    return records, sheet_info, auxiliary, duplicates, review_config


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
                    encoding="utf-8")


def build(root=ROOT):
    records, sheets, auxiliary, duplicates, review_config = extract(root)
    processed = root / "data/processed"
    out = root / "outputs/attachments"
    profiles = []
    for info in sheets:
        subset = [r for r in records if r["source_file"] == info["source_file"] and r["sheet"] == info["sheet"]]
        profiles.append({
            **info, "unique_names": len({r["name"] for r in subset}),
            "same_name_groups": len({r["duplicate_group"] for r in subset if r["duplicate_group"]}),
            "explicit_or_reviewed_province_records": sum(r["province"] is not None for r in subset),
            "province_label_counts": dict(sorted(Counter(r["province"] or "未解析" for r in subset).items())),
            "region_count_scope": "附件记录标签分布，非全国项目/容量分布；未独立核查的标签仍保留标记",
            "example_records": [r["record_id"] for r in subset[:2]],
            "partial_official_review_records": [r["record_id"] for r in subset if r["review_source_ids"]],
            "capacity_aggregation": None,
            "capacity_aggregation_reason": "年份、状态、项目范围、重复和单位未全部确认，禁止汇总为全国容量",
        })
    write_json(processed / "attachment_assets.json", {
        "schema_version": "1.0", "records": records, "source_sheets": sheets,
        "processing_code_sha256": sha256(Path(__file__)),
        "review_config_sha256": sha256(root / "data/attachment_reviews.json")})
    write_json(out / "sample_profiles.json", {"scope": "contest_attachment_samples_only", "profiles": profiles})
    write_json(out / "duplicate_review.json", {"groups": duplicates})
    write_json(out / "excluded_cells.json", {"cells": auxiliary})
    write_json(out / "issue_summary.json", {
        "records": len(records), "issues": dict(sorted(Counter(i for r in records for i in r["issues"]).items())),
        "partially_field_reviewed_records": len(review_config["reviews"]),
        "national_calibration_eligible": sum(r["national_calibration_eligible"] for r in records)})
    lines = ["# 题给附件标准化与用途报告", "",
             "原始四份Excel保持不变。输出逐行标准表保留原始单元格、公式与缓存值；未知年份和状态为空。",
             "以下计数是附件记录，不能解释为全国完整电站数或当前在运容量。", "",
             "|附件/工作表|具名记录|不同名称|同名组|有省份标签或官方修正的记录|",
             "|---|---:|---:|---:|---:|"]
    lines += [f"|{p['source_file']}/{p['sheet']}|{p['named_records']}|{p['unique_names']}|"
              f"{p['same_name_groups']}|{p['explicit_or_reviewed_province_records']}|" for p in profiles]
    lines += ["", f"共{len(records)}条具名记录。所有同名行保留；忽略空名称行"
              f"{sum(len(s['ignored_blank_name_rows']) for s in sheets)}行，其单元格内容另存审计。",
              "", "## 字段与单位", "",
              f"火电D列标题为KW，已逐行重算{sum(r['technology'] == 'thermal' for r in records)}条公式，"
              f"缓存不一致{sum('kw_cached_value_mismatch' in r['issues'] for r in records)}条。由明确KW列折算的MW只代表附件自身记载；"
              "不能由此确认项目范围、统计年份或当前运行状态。其他表的installedcapacity未标单位，不统一默认MW。",
              "水电D2为原表合计公式，独立留存、不并入任何电站容量。原表空格、点号与尾部格式行不作为观测。",
              "抽蓄装机功率与储能电量容量分开；附件没有时长和能量容量，不据此生成储能能量。",
              "", "## 同名记录", "",
              "|名称|原行号/记录|字段情况|处置|", "|---|---|---|---|"]
    lines += [f"|{g['name']}|{', '.join(str(x['row']) for x in g['records'])}|{g['kind']}|全部保留，待项目/机组身份核查|"
              for g in duplicates]
    lines += ["", "## 已有官方依据的逐字段核查", ""]
    for review in review_config["reviews"]:
        lines += [f"- {review['record_id']}：{review['reason']} 证据{review['source_id']}，定位：{review['locator']}。"]
    lines += ["", "核验只覆盖上述字段。核准容量不等于当年在运容量；投运表述不能反推精确投运日期。"
              "经纬度仅做数值范围检查，尚无坐标系与项目位置的完整验证，冲突记录隔离。",
              "", "## 附件已经进入的分析与后续使用", "",
              "|题给材料|本轮实际产物|对应题目/模型用途|限制|", "|---|---|---|---|",
              "|附件1火电|逐行资产档案、KW公式复核、同名组审计|问题一资产特征；W4投资/退役案例候选|燃料、年份、状态未明，不默认全为煤电|",
              "|附件2水电/抽蓄|分技术档案、省份冲突修正、官方项目案例|问题一协调性案例；W4水电与储能主体分类|抽蓄不计新增一次能源，能量容量待补|",
              "|附件3风电|16条档案、省份标签分布与案例清单|问题一说明样本地域覆盖；W4风电主体案例候选|地点和坐标未全核验，不能给全国主体权重|",
              "|附件4核电|17条档案、尾部空行排除与项目范围缺口|问题一核电资产案例清单；W4长寿命资产研究候选|期别、机组与统计年待核，不用17条代表全国|",
              "", "data/processed/attachment_assets.json为标准表；sample_profiles.json是可复算样本画像。"
              "本轮已实际使用四附件做样本结构与数据质量分析；模型初始化仍未接入，不能称已用于校准。",
              "", "## 下一步", "",
              "先选择可被官方证据覆盖的代表项目核查单位、期别及年份。全国模型容量由同口径全国官方序列提供，"
              "附件用于项目异质性、资产约束和区域案例，未经代表性设计不设置全国抽样权重。",
              "运行：run_attachments.ps1。字段含义及准入限制见docs/07_attachment_data_dictionary.md。"]
    out.mkdir(parents=True, exist_ok=True)
    (out / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {"records": len(records), "sheets": len(sheets), "duplicate_groups": len(duplicates),
            "official_field_reviews": len(review_config["reviews"])}


if __name__ == "__main__":
    print(json.dumps(build(), ensure_ascii=False))
