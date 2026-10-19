"""Extract the unchanged CEADs national 2014 sheet, retaining accounting scope.

This is a source transcription and diagnostic, never national model admission.
Reported net consumption is not gross final demand or a positive fuel supply.
"""
import csv
import json
import sys
from math import fsum, isfinite
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from energy_model.q2.io import csvout, dump, local_path, sha, verify_source_records


def classify(row):
    if row == 3: return 'final_total_control'
    if 4 <= row <= 50: return 'final_sector_detail'
    if row == 53: return 'conversion_total_control'
    if row in (54, 55, 56, 57, 58, 59, 61): return 'conversion_detail'
    if row == 63: return 'loss_separate_from_reported_total'
    if row == 65: return 'reported_net_consumption_control'
    return 'memo_not_additive'


def extract(sheet, review):
    values = list(sheet.iter_rows(values_only=True))
    if (sheet.max_row, sheet.max_column) != (70, 21):
        raise ValueError('National2014 source layout changed')
    if [r[0] for r in values] != review['row_labels']:
        raise ValueError('National2014 row labels changed')
    if list(values[0]) != review['carrier_labels'] or list(values[1]) != review['unit_labels']:
        raise ValueError('National2014 carrier/unit labels changed')
    rows, checks = [], []
    for row in range(3, 71):
        if not values[row - 1][0]:
            continue
        for col in range(2, 22):
            cell = sheet.cell(row, col)
            value = cell.value
            if cell.data_type == 'f' or (value is not None and
                    (type(value) not in (int, float) or not isfinite(value))):
                raise ValueError('Formula or non-numeric value requires separate review: ' + cell.coordinate)
            rows.append({'year':2014, 'geography':'China_national_CEADs_definition',
                'account':values[row-1][0].strip(), 'row_role':classify(row),
                'carrier':values[0][col-1], 'value':value,
                'unit_as_published':values[1][col-1], 'source_cell':cell.coordinate,
                'source_sheet':'2014', 'source_path':review['source_path'],
                'source_sha256':review['source_sha256'],
                'source_version':'10.6084/m9.figshare.5425546.v1',
                'status':'original_blank' if value is None else 'scoped_historical_reference',
                'transformation':'identity; original units, signs and missing values retained'})
    identities = [('final_sector_sum',3,list(range(4,51))),
                  ('conversion_detail_sum',53,[54,55,56,57,58,59,61]),
                  ('reported_total_excludes_loss',65,[3,53])]
    for name, total, parts in identities:
        for col in range(2,22):
            cells = [sheet.cell(row,col) for row in [total,*parts]]
            if any(type(c.value) not in (int,float) for c in cells):
                raise ValueError('Cannot assert identity with missing control or detail cells')
            residual = fsum(c.value for c in cells[1:]) - cells[0].value
            bound = .05 * len(cells)
            checks.append({'check':name,'carrier':values[0][col-1],
                'total_cell':cells[0].coordinate,'detail_cells':[c.coordinate for c in cells[1:]],
                'residual':residual,'rounding_bound':bound,'passed':abs(residual)<=bound+1e-8})
    if not all(c['passed'] for c in checks):
        raise ValueError('National2014 source arithmetic outside printed rounding bound')
    return rows, checks


def prepare(root=ROOT):
    import openpyxl
    review = json.loads((root/'data/catalog/q2_national_balance_review.json').read_text(encoding='utf-8'))
    if review['decision'] != 'scoped_2014_national_balance_reference_with_unresolved_discrepancies':
        raise ValueError('National balance reference review unresolved')
    if not review['checks'] or not all(c['passed'] and c['evidence'] for c in review['checks']):
        raise ValueError('National balance substantive review missing')
    for item in review['fingerprints']:
        if sha(local_path(root,item['path'])) != item['sha256']:
            raise ValueError('National balance reference changed: ' + item['path'])
    verify_source_records(root,review['source_records'])
    book = openpyxl.load_workbook(local_path(root,review['source_path']),read_only=True,data_only=False)
    try:
        if book.sheetnames != review['sheet_names']:
            raise ValueError('National inventory source coverage changed')
        rows, checks = extract(book['2014'],review)
    finally:
        book.close()
    cells = {r['source_cell']:r['value'] for r in rows}
    with (root/'data/processed/q2/nbs2014_accounts.csv').open(encoding='utf-8-sig',newline='') as stream:
        nbs = {(r['table'],r['account']):float(r['value']) for r in csv.DictReader(stream) if r['value']}
    comparisons = []
    for identity, addresses, nbs_key, note in [
        ('natural_gas_total',['R65'],('0909CH','total:gas'),'Reported net consumption excludes loss; source/method difference remains unresolved.'),
        ('natural_gas_transport',['R46'],('0909CH','transport:gas'),'Sector labels do not resolve the numerical discrepancy.'),
        ('natural_gas_household',['R49','R50'],('0909CH','household:gas'),'Only this local comparison agrees within workbook rounding.'),
        ('raw_coal_power_input',['B54'],('0905CH','power_input'),'Different input totals retained; not rescaled or averaged.'),
        ('final_electricity',['T3'],('0906CH','final'),'Equal control total does not validate other carriers.')]:
        left, right = fsum(cells[c] for c in addresses), nbs[nbs_key]
        comparisons.append({'comparison':identity,'ceads_cells':addresses,'ceads_value':left,
            'nbs_table':nbs_key[0],'nbs_account':nbs_key[1],'nbs_value':right,
            'difference_ceads_minus_nbs':left-right,'interpretation':note,
            'status':'agreement_within_rounding' if abs(left-right)<=.05*len(addresses)+.005 else 'unresolved_difference'})
    out = root/'data/processed/q2'
    csvout(out/'national_balance_2014_reference.csv',rows)
    csvout(out/'national_balance_2014_comparisons.csv',comparisons)
    summary = {'year':2014,'rows':len(rows),'checks':checks,
        'all_arithmetic_checks_passed':True,'national_initialization_admitted':False,
        'arithmetic_is_not_method_or_vintage_reconciliation':True,
        'formula_cells_accepted':0,'original_blank_count':sum(r['value'] is None for r in rows),
        'source_sha256':review['source_sha256'],
        'outputs':{n:sha(out/n) for n in ['national_balance_2014_reference.csv','national_balance_2014_comparisons.csv']}}
    dump(out/'national_balance_2014_checks.json',summary)
    return {'rows':len(rows),'identities':len(checks),'all_arithmetic_checks_passed':True,
            'unresolved_cross_source_differences':sum(r['status']=='unresolved_difference' for r in comparisons),
            'national_initialization_admitted':False}


if __name__ == '__main__': print(json.dumps(prepare(),ensure_ascii=False))
