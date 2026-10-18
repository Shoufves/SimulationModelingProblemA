"""Rebuild Q1 review and additions OFFLINE from the locked P1 source materials.

This is an explicit research preparation step, not run by `run.py q1`.
Fingerprints are recorded only after semantic checks against source cells,
source transcription and reviewed PDF/image locations pass.
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from energy_model.q1 import read_csv, read_json, write_csv, write_json, sha, Q1Error, review_snapshot


def main():
    import math
    import re
    import openpyxl
    from pypdf import PdfReader
    reviews = []
    def check(name, ok, detail):
        if not ok:
            raise Q1Error(name + ": " + str(detail))
        reviews.append(dict(check=name, passed=True, detail=detail))
    def equal(a, b):
        return math.isclose(float(a), float(b), rel_tol=1e-11, abs_tol=1e-8)
    def rows(name):
        return read_csv(ROOT / ("data/processed/" + name + ".csv"))

    # Images C09-01/02/03/04/05/06/14/15/19 were visually read by mmj on
    # 2026-10-18. Historical values retain NBS units and table boundaries.
    tr = read_json(ROOT / "data/extracted/nbs_verified_transcription.json")
    expected = {}
    for group in ("energy_annual", "capacity_annual", "products_annual"):
        table = tr[group]
        for line in table["rows"]:
            for var, value in zip(table["columns"][1:], line[1:]):
                if value is not None:
                    expected[(int(line[0]), var)] = value
    nbs = [r for r in rows("unified/nbs_history")
           if r["family"] in ("energy", "capacity") and 2015 <= int(r["year"]) <= 2024]
    products = [r for r in rows("nbs_products_annual") if 2015 <= int(r["year"]) <= 2024]
    check("NBS_unique_view_and_products_vs_reviewed_transcription",
          all(equal(r["value"], expected[(int(r["year"]), r["variable"])])
              for r in nbs + products), dict(rows=len(nbs+products),
              image_review="C09-01/02/15/19; 2015-2024 rows and notes",
              caution="same-source concordance, not independent validation"))
    balances = rows("nbs_balance_checkpoints")
    b = tr["balance_checkpoints"]
    check("NBS_balance_cells",
          all(equal(r["value"], b["tables"][r["source_id"]]["rows"][r["variable"]][
              b["years"].index(int(r["year"]))]) for r in balances),
          "164 cells; C09-03/04/05/06 selected-year columns; nested rows not additive")
    d = tr["province_demand"]
    demand_expected = {(int(y), line[0]): value for line in d["rows"]
                       for y, value in zip(d["years"], line[1:])}
    demand = rows("nbs_province_demand")
    check("NBS_province_demand_cells",
          all(equal(r["value"], demand_expected[(int(r["year"]), r["region"])]) for r in demand),
          "124 cells; C09-14; provincial use rounded to integer 10^8 kWh")
    # New transcription is an incremental interpretation of an EXISTING base image.
    # Blank cells are explicit nulls, not zeros or inferred production.
    provincial = [
        ("北京", None, 1.20, 455.64, 9.44), ("天津", None, 45.42, 879.95, .14),
        ("河北", .44, 5.05, 4334.58, 65.38), ("山西", 12.79, 167.39, 4629.49, 48.33),
        ("内蒙古", 13.00, 309.99, 8344.00, 55.44),
        ("辽宁", .31, 8.13, 2379.37, 77.71), ("吉林", .09, 18.27, 1179.62, 124.78),
        ("黑龙江", .54, 53.85, 1372.34, 47.48), ("上海", None, 50.25, 1078.02, None),
        ("江苏", .09, 1.38, 6807.37, 32.73), ("浙江", None, None, 4983.55, 263.31),
        ("安徽", 1.06, 2.54, 3863.46, 94.52), ("福建", .04, None, 3584.41, 445.19),
        ("江西", .03, .09, 1963.38, 188.82), ("山东", .87, 8.60, 6773.25, 54.53),
        ("河南", 1.04, 4.27, 3832.45, 146.08), ("湖北", .01, 2.11, 3498.65, 1305.91),
        ("湖南", .09, .20, 1900.11, 509.93), ("广东", None, 126.69, 7354.51, 429.57),
        ("广西", .03, .33, 2602.85, 614.85), ("海南", None, 35.64, 515.29, 30.77),
        ("重庆", None, 98.92, 1210.10, 218.40), ("四川", .22, 656.68, 5307.27, 4045.43),
        ("贵州", 1.46, 11.91, 2556.10, 616.60), ("云南", .67, .20, 4646.34, 3304.33),
        ("西藏", None, None, 171.63, 132.71), ("陕西", 7.82, 359.21, 3374.06, 88.85),
        ("甘肃", .67, 9.12, 2271.50, 396.01), ("青海", .08, 59.46, 1153.50, 507.27),
        ("宁夏", 1.04, 4.86, 2368.02, 19.84), ("新疆", 5.43, 422.75, 5478.02, 382.46)]
    out = ROOT / "data/processed/q1"
    out.mkdir(parents=True, exist_ok=True)
    fields = ["year", "region", "variable", "value", "unit", "source_id",
              "source_locator", "status", "boundary", "missing_reason"]
    records = []
    for region, *values in provincial:
        for (var, unit), value in zip([
            ("raw_coal_100Mt", "10^8 t"), ("natural_gas_100Mm3", "10^8 m3"),
            ("generation_100GWh", "10^8 kWh"), ("hydro_generation_100GWh", "10^8 kWh")], values):
            records.append(dict(year=2024, region=region, variable=var, value=value,
                unit=unit, source_id="NBS_2025_C09-19",
                source_locator=f"table 9-19; 2024 provincial section; row={region}; column={var}",
                status="reviewed_transcription" if value is not None else "missing_source_cell",
                boundary="China mainland provinces; source product table; no inferred interprovincial flows",
                missing_reason="" if value is not None else "原表空白，原因未说明，不填零"))
    write_csv(out / "province_products_2024.csv", records, fields)
    reviews.append(dict(check="new_provincial_transcription_visual_review", passed=True,
        detail="mmj/2026-10-18 viewed C09-19 image: 31 province names and four columns; "
               "blank Shanghai hydro retained; C09-14 2024 demand column independently reread. "
               "One transcription, awaiting teammate review; not an independent observation."))

    raw_macro = read_json(ROOT / "data/raw/batch1_national/wdi_china_macro_2005_2025.json")
    macro_expected = {(int(r["date"]), r["indicator"]["id"]): r["value"] for r in raw_macro[1]
                      if r["countryiso3code"] == "CHN"}
    macro = [r for r in rows("china_macro_annual")
             if r["variable"] == "NY.GDP.MKTP.KN" and 2015 <= int(r["year"]) <= 2024]
    check("WDI_constant_GDP_raw_JSON", all(equal(r["value"], macro_expected[
        (int(r["year"]), r["variable"])]) for r in macro),
        dict(rows=len(macro), version=raw_macro[0]["lastupdated"],
             scope="source series constant LCU base; not asserted to be 2020-price GDP"))
    raw_co2 = {int(r["Year"]): r for r in read_csv(
        ROOT / "data/raw/batch1_national/GCB2025v15_MtCO2_flat.csv")
        if r["ISO 3166-1 alpha-3"] == "CHN" and 2015 <= int(r["Year"]) <= 2024}
    co2 = [r for r in rows("china_co2_gcb_annual") if 2015 <= int(r["year"]) <= 2024]
    check("GCB_raw_country_component_cells", all(equal(r["value"],
        raw_co2[int(r["year"])][r["variable"]]) for r in co2),
        dict(rows=len(co2), scope="GCB2025v15 territorial fossil and industry; no removals"))

    ei_path = ROOT / "data/raw/lb/trade_alternatives/ei_2025_data.xlsx"
    book = openpyxl.load_workbook(ei_path, data_only=True, read_only=True)
    for file in ("pipeline_gas_china_2024", "lng_china_2024", "coal_china_2024_EJ"):
        table = rows("lb/trade_alternatives/" + file)
        # Boundaries and headings belong to the specified sheet, not just cell values.
        sheet = book[table[0]["sheet"]]
        coordinates=[openpyxl.utils.cell.coordinate_to_tuple(r["cell"]) for r in table]
        grid = list(sheet.iter_rows(max_row=max(r for r,c in coordinates),
                    max_col=max(c for r,c in coordinates),values_only=True))
        def cell_value(cell):
            row, col = openpyxl.utils.cell.coordinate_to_tuple(cell)
            return grid[row-1][col-1]
        check("EI_cells_" + file, all(equal(r["value"], cell_value(r["cell"])) for r in table),
              dict(rows=len(table), sheet=sheet.title, cells=[r["cell"] for r in table],
                   scope="2024 China imports; source bins include aggregates; GCV 40 MJ/m3 gas or EJ coal"))
    book.close()
    wb = openpyxl.load_workbook(ROOT / "data/raw/batch1_national/wb_commodity_prices_stable.xlsx",
                               data_only=True, read_only=True)
    sheet = wb["Annual Prices (Nominal)"]
    grid = list(sheet.values)
    headers = {str(v): i for i, v in enumerate(grid[6]) if v is not None}
    annual = {int(line[0]): line for line in grid[7:] if isinstance(line[0], (int, float))}
    prices = [r for r in rows("commodity_price_references")
              if r["status"] == "archived_reference" and 2015 <= int(r["year"]) <= 2024]
    check("WB_archived_workbook_cells", all(equal(r["value"],
          annual[int(r["year"])][headers[r["variable"]]]) for r in prices),
          dict(rows=len(prices), sheet=sheet.title, header_row=7,
               coverage="2015-2020 only; quarantined edition excluded, not bridged"))
    wb.close()

    irena = PdfReader(ROOT / "data/raw/batch3_scenarios/irena_power_costs_2024.pdf")
    page = lambda n: irena.pages[n-1].extract_text()
    # Actual source table rows, not tokens in a derived CSV.
    assertions = [(71, r"(?m)^China\s+1\s+489\s+1\s+765\s+2\s+066\s+626\s+856\s+1\s+315"),
                  (95, r"(?m)^China\s+1\s+523\s+1\s+623\s+2\s+767\s+438\s+591\s+867"),
                  (119, r"(?m)^China\s+3\s+307\s+5\s+268\s+5\s+852\s+1\s+434\s+1\s+520\s+1\s+869"),
                  (76, r"China\s+25\s+35\s+33"),
                  (123, r"China\s+.*?37"),
                  (209, r"China\s+29"), (212, r"China\s+59"),
                  (202, r"Wind power\s+25"),
                  (84, r"China\s+.*?0\.023\s+0\.029\s+0\.048"),
                  (106, r"China\s+.*?0\.028\s+0\.033\s+0\.052"),
                  (126, r"China\s+.*?0\.046\s+0\.056\s+0\.063")]
    for p, pattern in assertions:
        check(f"IRENA_PDF_p{p}_table_row", bool(re.search(pattern, page(p), re.S)),
              "China 2024 row / Annex method; source physical PDF page, not printed offset")
    cost_expected = {
        "onshore_wind_installed_cost_weighted_average":856,
        "onshore_wind_installed_cost_p05":626,"onshore_wind_installed_cost_p95":1315,
        "utility_solar_PV_installed_cost_weighted_average":591,
        "utility_solar_PV_installed_cost_p05":438,"utility_solar_PV_installed_cost_p95":867,
        "offshore_wind_installed_cost_weighted_average":1520,
        "offshore_wind_installed_cost_p05":1434,"offshore_wind_installed_cost_p95":1869,
        "onshore_wind_LCOE":.029,"utility_solar_PV_LCOE":.033,"offshore_wind_LCOE":.056}
    cost_rows = [r for r in rows("technology_accounting_parameters")
                 if r["source_id"]=="IRENA_COST_2024_PDF" and int(r["year"])==2024]
    check("cost_table_vs_reviewed_PDF_rows",
          len(cost_rows)==len(cost_expected) and all(equal(r["value"],cost_expected[r["variable"]])
          for r in cost_rows), "12 adopted China 2024 CAPEX/quantile/LCOE values checked")
    cost = []
    for technology, om, cf, ompage, cfpage, omtable, cftable in [
        ("onshore_wind", 29, .33, 209, 76, "A3", "2.2"),
        ("offshore_wind", 59, .37, 212, 123, "A6", "4.3")]:
        for variable, value, unit, p, table in [
            ("fixed_om", om, "2024 USD/kW/year", ompage, omtable),
            ("new_project_capacity_factor", cf, "fraction", cfpage, cftable),
            ("economic_life", 25, "year", 202, "A1")]:
            cost.append(dict(year=2024, region="China", variable=technology+"_"+variable,
                value=value, unit=unit, source_id="IRENA_COST_2024_PDF",
                source_locator=f"PDF p{p}; Table {table}; "
                               + ("Wind power economic-life assumption" if variable=="economic_life"
                                  else "China 2024"),
                status="reviewed_published_assumption" if variable != "new_project_capacity_factor"
                       else "reviewed_new_project_parameter",
                boundary="New commissioned projects; CF generally ex ante; country averages not matched project data; "
                         "O&M is a published assumption; not national operating fleet; plant-gate unfirmed electricity",
                missing_reason=""))
    write_csv(out / "cost_method_parameters.csv", cost, fields)

    # Existing locked-base PDF: PDF p21 / printed p011, read with Figure 1.
    # Preserve paragraph and figure precision separately; neither is initial SOC.
    storage_pdf = PdfReader(ROOT / "data/raw/lb/base/reports/storage_2025.pdf")
    storage_text = re.sub(r"\s+", "", storage_pdf.pages[20].extract_text())
    storage_heading = re.sub(r"\s+", "", storage_pdf.pages[19].extract_text())
    check("storage_2024_paragraph_and_unit_review",
          "截至2024年底" in storage_heading and
          all(token in storage_text for token in ("7376万千瓦/1.68亿千瓦时", "2.3小时")),
          "PDF p20-21 / printed pp010-011; 2024 commissioned cumulative stock, not 2025. "
          "mmj visually checked p21 Figure 1: 7376 (10^4 kW), 16820 (10^4 kWh), 2.3 h; "
          "paragraph energy 1.68*10^8 kWh retained separately from figure precision. Peer review pending.")
    storage_records = []
    for variable, value, unit, locator in [
        ("new_storage_power", 7376, "10^4 kW", "opening paragraph and Figure 1; 2024 power"),
        ("new_storage_energy_paragraph", 1.68, "10^8 kWh", "opening paragraph; 2024 energy"),
        ("new_storage_energy_figure", 16820, "10^4 kWh", "Figure 1; 2024 energy bar"),
        ("new_storage_average_duration", 2.3, "h", "opening paragraph and Figure 1; 2024 duration")]:
        storage_records.append(dict(year=2024, region="China", variable=variable, value=value, unit=unit,
            source_id="NEA_STORAGE_2025", source_locator="PDF p21 / printed p011; "+locator,
            status="reviewed_published_rounded",
            boundary="End-2024 commissioned new-type storage; excludes pumped hydro; nameplate not SOC; "
                     "paragraph/figure energy precision retained separately; model-region mapping pending",
            missing_reason=""))
    write_csv(out / "storage_2024.csv", storage_records, fields)

    # Review all adopted snapshot fields against their source paragraphs.
    gas = PdfReader(ROOT / "data/raw/batch2_mechanisms/nea_natural_gas_2025.pdf")
    h2 = PdfReader(ROOT / "data/raw/batch3_scenarios/nea_hydrogen_2025.pdf")
    snapshots = rows("fuel_storage_hydrogen_snapshots")
    texts = {
        "NEA_NATURAL_GAS_2025": "\n".join(gas.pages[n-1].extract_text() for n in (10,11,12,13)),
        "NEA_HYDROGEN_2025": "\n".join(h2.pages[n-1].extract_text() for n in (16,17)),
        "NDRC_GAS_APPARENT_2024": (ROOT / "data/raw/batch1_national/ndrc_gas_apparent_2024.html").read_text(encoding="utf-8"),
        "NEA_STORAGE_2025": (ROOT / "data/raw/batch2_mechanisms/nea_storage_2025.html").read_text(encoding="utf-8")}
    from lxml import html
    for source in ("NDRC_GAS_APPARENT_2024","NEA_STORAGE_2025"):
        tree=html.fromstring(texts[source])
        for node in tree.xpath("//script|//style"):
            node.drop_tree()
        texts[source]=tree.text_content()
    def numeric_token(value, text):
        value=float(value)
        token=re.escape(str(int(abs(value))) if value.is_integer() else format(abs(value),".15g"))
        if value.is_integer():
            token+=r"(?:\.0+)?"
        clean=re.sub(r"\s+","",text).replace(",","")
        prefix=r"(?:-|下降)" if value<0 else r"(?<![\d.])"
        return bool(re.search(prefix+token+r"(?![\d.])",clean))
    selected_snapshots = [r for r in snapshots if int(r["year"]) == 2024 or
                          r["variable"] in ("new_storage_power", "new_storage_energy", "new_storage_average_duration")]
    storage_expected = {"new_storage_power":136,"new_storage_energy":351,
                        "new_storage_average_duration":2.58}
    check("storage_original_unit_conversion",
          "1.36亿千瓦/3.51亿千瓦时" in texts["NEA_STORAGE_2025"] and
          "2.58小时" in texts["NEA_STORAGE_2025"],
          "1.36*100 GW, 3.51*100 GWh; 2025 installed source paragraph read with year")
    snapshot_mismatches = [r["variable"] for r in selected_snapshots if not (
        equal(r["value"],storage_expected[r["variable"]]) if r["source_id"]=="NEA_STORAGE_2025"
        else numeric_token(r["value"],texts[r["source_id"]]))]
    check("snapshot_source_paragraphs", not snapshot_mismatches, dict(rows=len(selected_snapshots),
          mismatches=snapshot_mismatches,
          locators="gas PDF pp10-13; hydrogen pp16-17; NDRC annual paragraph; NEA storage installed statement",
          qualifiers="H2 total EXCEEDS 3650, categories approximate; renewable H2 capacity not output; storage is 2025"))
    for name, capture in [("energy_official_communique_2024", "S13"),
                          ("renewable_operation_official_2024", "S25")]:
        obs = read_json(ROOT / f"data/observations/{name}.json")
        text = (ROOT / f"data/raw/official/{capture}_selected_capture.txt").read_text(encoding="utf-8")
        mismatches=[r["indicator_id"] for r in obs["observations"] if not numeric_token(r["value"],text)]
        check(capture + "_reviewed_capture", not mismatches, dict(rows=len(obs["observations"]),mismatches=mismatches,
              scope="Existing review/capture read with notes; adopted uses stay within original contract"))
    review_path = ROOT / "data/catalog/q1_evidence_review.json"
    review = dict(reviewer="mmj", review_date="2026-10-18",
        base_sha="7a65d068d19eea2f7dd8161a618a3ad8d66f98b8",
        scope="Q1 descriptive research only; source agreement does not grant Q2 model admission",
        checks=reviews,
        manual_review="NBS image rows, headings and notes viewed in this session; IRENA tables/method "
                      "and gas/H2 paragraphs read from stored original PDFs. Two teammate reviews pending.",
        exclusions=["data/restricted/**", "quarantined WB edition and lb price files",
                    "candidate transmission network", "regional storage shares not mapped to Q2 regions"])

    registry = {s["source_key"]: s for s in read_json(ROOT / "data/catalog/sources_unified.json")}
    ei = next(k for k in registry if k.endswith(":EI_DATA_2025") and "alternatives" in k)
    specs = [
        ("nbs", "data/processed/unified/nbs_history.csv", ["yx:NBS_2025_C09-01","yx:NBS_2025_C09-02","yx:NBS_2025_C09-15"],
         "2015-2024 energy/capacity only; exclude both GDP families", "NBS2025 same-table history; end-year capacity; primary electricity/other not relabelled"),
        ("products", "data/processed/nbs_products_annual.csv", ["yx:NBS_2025_C09-19"], "2015-2024", "National product output; 10^8 kWh = 0.1 TWh"),
        ("balance", "data/processed/nbs_balance_checkpoints.csv", ["yx:NBS_2025_C09-03","yx:NBS_2025_C09-04","yx:NBS_2025_C09-05","yx:NBS_2025_C09-06"],
         "2015/2020/2022/2023 only", "Same-table ledgers; oil includes products/bunkering; nested transformation not added"),
        ("demand", "data/processed/nbs_province_demand.csv", ["yx:NBS_2025_C09-14"], "2015/2020/2023/2024", "31 provinces; integer-rounded use; own annual denominator"),
        ("province", "data/processed/q1/province_products_2024.csv", ["yx:NBS_2025_C09-19"], "2024", "New visual transcription from locked existing image; blanks not zero; candidate increment pending peer review"),
        ("macro", "data/processed/china_macro_annual.csv", ["yx:WDI_CHINA_MACRO"], "2015-2024 constant GDP only", "2026-10-08 version; constant LCU source base; CHN boundary bridge to mainland needs review for Q2"),
        ("co2", "data/processed/china_co2_gcb_annual.csv", ["yx:GCB_GCB2025v15_MtCO2_flat"], "2015-2024", "GCB2025v15 CHN territorial fossil/industry; country bridge to mainland pending Q2; not net GHG"),
        ("snapshot", "data/processed/fuel_storage_hydrogen_snapshots.csv", ["yx:NEA_NATURAL_GAS_2025","yx:NEA_HYDROGEN_2025","yx:NDRC_GAS_APPARENT_2024","yx:NEA_STORAGE_2025"],
         "2024 gas/H2; 2025 national storage power/energy/duration separately", "Gas definitions not bridged to EI; H2 exceeds/approx qualifiers; added storage not inventory"),
        ("cost", "data/processed/technology_accounting_parameters.csv", ["yx:IRENA_COST_2024_PDF"], "2024 China IRENA only", "CAPEX/quantiles/LCOE; IPCC and 2022 market price excluded; LCOE not additive to capex"),
        ("cost_method", "data/processed/q1/cost_method_parameters.csv", ["yx:IRENA_COST_2024_PDF"], "2024 new wind projects", "25y; source CF and O&M; common 3/5/7% real WACC is analysis assumption, not observed China financing"),
        ("storage_2024", "data/processed/q1/storage_2024.csv", ["lb/base:NEA_STORAGE_2025"],
         "2024 cumulative commissioned new-type storage, from report published 2025",
         "Existing locked-base PDF p21; paragraph/figure energy precision separate; P/E not initial SOC; peer confirmation pending"),
        ("price", "data/processed/commodity_price_references.csv", ["yx:WB_PINK_SHEET_STABLE"], "2015-2024 requested; only 2015-2020 archived rows adopted", "Nominal international benchmarks; no quarantined edition, no China contract-price inference"),
        ("pipeline", "data/processed/lb/trade_alternatives/pipeline_gas_china_2024.csv", [ei], "2024", "China pipeline imports; bcm GCV40MJ/m3; bins not invented countries"),
        ("lng", "data/processed/lb/trade_alternatives/lng_china_2024.csv", [ei], "2024", "China LNG imports same EI standard; no mixing with domestic gas"),
        ("coal_trade", "data/processed/lb/trade_alternatives/coal_china_2024_EJ.csv", [ei], "2024", "EJ imports; grouped geography gives country HHI bounds"),
        ("communique", "data/observations/energy_official_communique_2024.json", ["mmj:S13"], "2024 independent release", "Selected crude oil/trade/generation/clean share/growth; no unbridged history splice"),
        ("renewables", "data/observations/renewable_operation_official_2024.json", ["mmj:S25"], "2024 independent release", "15 original scoped observations; version comparison only when NBS differs")]
    inputs, sources = [], {}
    for key, path, keys, coverage, boundary in specs:
        for source_key in keys:
            source = registry[source_key]
            p = source["canonical_path"]
            if not p:
                p = source["evidence_files"][0]["path"]
            sources[source_key] = dict(path=p, sha256=sha(ROOT/p), url=source["url"],
                version="Locked base source edition; see source_key and per-record locator",
                content_review="q1_evidence_review.json", use="Q1 scoped research only")
        inputs.append(dict(id=key, path=path, sha256=sha(ROOT/path), source_keys=keys,
            coverage=coverage, boundary=boundary,
            evidence_status="reviewed_Q1_only_pending_two_peer_reviews",
            allowed_use="Q1 descriptive research, diagnostics and Q2 candidate handoff",
            blocked_use="automatic national ABM initialization/calibration/forecast",
            locator="Per-row source_locator; JSON observations locator; trade sheet/cell"))
    # Capture this only after the substantive checks above; refreshing a manifest
    # fingerprint alone cannot reuse an older review for changed data or use.
    review["reviewed_snapshot"] = review_snapshot(inputs, sources)
    write_json(review_path, review)
    manifest = dict(schema_version=1, candidate="p1/mmj", base_sha="7a65d068d19eea2f7dd8161a618a3ad8d66f98b8",
        reviewed_on="2026-10-18", observation_window=[2015,2024],
        common_input_change=False,
        incremental_extraction="Three processed tables use originals already in locked base; province, cost method and 2024 storage are candidate additions pending Issue/peer confirmation, not replacement shared baseline",
        assumptions=dict(real_wacc=[.03,.05,.07], common_life_years=25,
                         cost_service="unfirmed plant-gate electricity; constant annual output; 8760 hours normalization"),
        evidence_review=dict(path="data/catalog/q1_evidence_review.json", sha256=sha(review_path)),
        inputs=inputs, sources=sources)
    write_json(ROOT/"workflow/q1_inputs.json", manifest)
    # Update only new rows in the shared catalog, preserving every existing row.
    catalog_path = ROOT / "data/catalog/datasets_unified.csv"
    newpaths = {s[1] for s in specs if "/q1/" in s[1]}
    original_lines=catalog_path.read_bytes().splitlines(keepends=True)
    retained_lines=[line for line in original_lines
                    if not any(line.startswith(p.encode("utf-8")+b",") for p in newpaths)]
    additions=[]
    for item in inputs:
        if item["path"] in newpaths:
            additions.append(dict(path=item["path"], final_refs="", status="reviewed_Q1_candidate",
                use=item["allowed_use"], boundary=item["boundary"], sha256=item["sha256"]))
    # Existing catalog rows retain their original bytes and newline convention.
    import csv,io
    line_ending="\r\n" if any(line.endswith(b"\r\n") for line in retained_lines) else "\n"
    buffer=io.StringIO(newline="")
    writer=csv.DictWriter(buffer,["path","final_refs","status","use","boundary","sha256"],
                          lineterminator=line_ending)
    writer.writerows(additions)
    catalog_path.write_bytes(b"".join(retained_lines)+buffer.getvalue().encode("utf-8"))
    print(f"Q1 preparation: {len(reviews)} substantive checks; {len(inputs)} inputs; no network.")


if __name__ == "__main__":
    main()
