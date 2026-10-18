import json,csv,re
from pathlib import Path
from openpyxl import load_workbook
from html.parser import HTMLParser
P=Path(__file__).resolve().parent; D=P/'data';D.mkdir(exist_ok=True)
def out(name,rows):
    if not rows:return
    with (D/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,list(rows[0]));w.writeheader();w.writerows(rows)
routes={3:'Malacca',5:'Hormuz',7:'Suez_SUMED',9:'Danish_Straits',11:'Dardanelles',14:'Cape_of_Good_Hope'}
rows=[]
for n,route in routes.items():
    p=P/'raw'/f'eia_chokepoints_F{n}.xlsx';w=load_workbook(p,read_only=True,data_only=True)
    for ws in w:
        data=list(ws.values);title=str(data[0][0]);unit=data[1][0]
        direction='origin' if 'origin' in title.lower() else 'destination' if 'destination' in title.lower() else 'other'
        for ri,row in enumerate(data[2:],3):
            if not row[0] or not any(isinstance(v,(float,int)) for v in row[1:7]):continue
            for ci,period in enumerate(data[1][1:7],1):
                value=row[ci] if len(row)>ci else None
                if not isinstance(value,(int,float)):continue
                rows.append(dict(route=route,direction=direction,country_or_group=row[0],period=str(period),period_type='half_year' if period=='1H25' else 'year',value=value,unit=unit,commodity='crude oil and condensate',source_id=f'EIA_FIGURE_{n}',source_sheet=ws.title,source_cell=f'{chr(65+ci)}{ri}',note='Flows may cross multiple routes; do not sum routes as independent imports. Geography label retained from source.'))
    w.close()
out('route_crude_origin_destination_2020_1H2025.csv',rows)
china=[r for r in rows if r['direction']=='destination' and r['country_or_group']=='China']
out('china_destination_route_flows_2020_1H2025.csv',china)

class Tables(HTMLParser):
    def __init__(self):super().__init__();self.tables=[];self.depth=0;self.t=None;self.row=None;self.cell=None;self.caption=False
    def handle_starttag(self,tag,attrs):
        if tag=='table':
            self.depth+=1
            if self.depth==1:self.t={'caption':'','rows':[]}
        if self.depth==1:
            if tag=='tr':self.row=[]
            if tag in ('td','th'):self.cell=[]
            if tag=='caption':self.caption=True
    def handle_data(self,data):
        if self.depth==1:
            if self.cell is not None:self.cell.append(data)
            if self.caption:self.t['caption']+=data
    def handle_endtag(self,tag):
        if self.depth==1:
            if tag in ('td','th') and self.cell is not None:
                if self.row is not None:self.row.append(re.sub(r'\s+',' ',''.join(self.cell)).strip())
                self.cell=None
            if tag=='tr' and self.row is not None:self.t['rows'].append(self.row);self.row=None
            if tag=='caption':self.caption=False
            if tag=='table':self.tables.append(self.t)
        if tag=='table':self.depth-=1
t=Tables();t.feed((P/'raw/eia_security_aug2026.html').read_text(encoding='utf-8'))
(D/'eia_security_tables_raw.json').write_text(json.dumps(t.tables,ensure_ascii=False,indent=2),encoding='utf-8')
security=[]
for ti,tab in enumerate(t.tables,1):
    header=next((r for r in tab['rows'] if sum(bool(re.fullmatch(r'[1-4]Q\d{2}',c)) for c in r)>=3),None)
    if not header:continue
    periods=[c for c in header if re.fullmatch(r'[1-4]Q\d{2}',c)]
    for ri,row in enumerate(tab['rows'],1):
        if len(row)!=len(periods)+1 or row==header:continue
        unit='million barrels' if ti==1 else 'billion cubic feet per day' if 'LNG' in row[0] else 'million barrels per day'
        for period,raw in zip(periods,row[1:]):
            try:value=float(raw.replace(',',''))
            except ValueError:value=''
            security.append(dict(table=ti,caption=tab['caption'].strip(),series=row[0],period=period,value=value,raw_value=raw,unit=unit,source_id='EIA_SECURITY_AUG2026',source_row=ri,data_type='EIA_estimate',note='Third-party estimates, not China official inventories; dash retained as missing.'))
out('eia_security_quarterly_2025_2026.csv',security)
checks=[]
for table in (2,3):
    matches=[r for r in security if r['table']==table and 'Malacca' in r['series'] and 'LNG' not in r['series'] and r['period']=='4Q25']
    checks.extend(matches)
(D/'quality_notes.json').write_text(json.dumps({'route_rows':len(rows),'china_route_rows':len(china),'security_rows':len(security),'malacca_4Q25_cross_table_review':checks,'warnings':['Same cargo may cross Hormuz and Malacca; avoid double-counting.','Origin and destination marginal tables are not a bilateral origin-destination matrix.','EIA China destination definition needs cross-check with China mainland customs.','EIA strategic inventory estimates do not establish officially usable reserve volumes.','STEO workbook is raw reference: forecasts, estimates and placeholder zeros require field-level verification.']},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'route_rows':len(rows),'china_rows':len(china),'security_rows':len(security),'security_tables':len(t.tables),'malacca_review':checks},ensure_ascii=False))
