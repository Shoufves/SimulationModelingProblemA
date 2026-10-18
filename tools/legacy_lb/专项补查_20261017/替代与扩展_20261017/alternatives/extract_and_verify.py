"""Extract public aggregate trade tables; preserve the original workbook unchanged."""
from pathlib import Path
import csv, json, hashlib, math
import openpyxl

P = Path(__file__).resolve().parent
D = P / 'data'
D.mkdir(exist_ok=True)
w = openpyxl.load_workbook(P/'raw/ei_2025_data.xlsx', read_only=True, data_only=True)
checks = []
def save(name, rows):
    with (D/name).open('w', encoding='utf-8-sig', newline='') as f:
        wr=csv.DictWriter(f,fieldnames=list(rows[0]));wr.writeheader();wr.writerows(rows)

for name, rowno, header, endcol, unit, out in [
    ('Gas trade 2024 - pipeline',27,3,23,'billion cubic metres; GCV 40 MJ/m3','pipeline_gas_china_2024.csv'),
    ('Gas trade 2024 - LNG',28,4,24,'billion cubic metres; GCV 40 MJ/m3','lng_china_2024.csv')]:
    s=w[name]
    matrix=list(s.iter_rows(max_row=43,max_col=endcol,values_only=True))
    assert matrix[rowno-1][0]=='China'
    total=matrix[rowno-1][endcol-1]
    rows=[]
    for col in range(2,endcol+1):
        origin=str(matrix[header-1][col-1]).strip()
        if header==4:
            top=matrix[2][col-1]
            if top is not None: origin=str(top).strip()+' '+origin
        cell=openpyxl.utils.get_column_letter(col)+str(rowno)
        value=matrix[rowno-1][col-1]
        rows.append(dict(year=2024,destination='China',origin=origin,value=value,unit=unit,
                         share_within_source=(value/total if col<endcol else ''),
                         source_id='EI_DATA_2025',sheet=name,cell=cell,
                         row_type='total' if col==endcol else 'origin'))
    subtotal=sum(r['value'] for r in rows if r['row_type']=='origin')
    assert math.isclose(subtotal,total,rel_tol=1e-10)
    checks.append(dict(table=out,rows=len(rows),sum=subtotal,total=total,passed=True))
    save(out,rows)

s=w['Coal trade movements 2024'];m=list(s.iter_rows(max_row=22,max_col=15,values_only=True))
assert m[2][9]=='China'
rows=[dict(year=2024,destination='China',origin=m[r-1][0],value=m[r-1][9],unit='exajoules',
           source_id='EI_DATA_2025',sheet=s.title,cell='J'+str(r),row_type='total' if r==19 else 'origin') for r in range(5,20)]
assert math.isclose(sum(r['value'] for r in rows[:-1]),rows[-1]['value'],rel_tol=1e-10)
save('coal_china_2024_EJ.csv',rows)
checks.append(dict(table='coal_china_2024_EJ.csv',rows=len(rows),passed=True))
countries={'印度尼西亚':[22025,24088,21145],'澳大利亚':[5247,8324,7747],'俄罗斯':[10213,9507,8889],
           '蒙古':[6997,8291,8896],'加拿大':[886,937,1130],'菲律宾':[498,772,655],
           '美国':[683,1213,291],'南非':[168,68,53],'哥伦比亚':[438,751,40]}
save('coal_origins_2023_2025_CCTD.csv',[dict(origin=c,year=y,quantity=v,unit='万吨',source_id='COAL_ORIGIN_CHART_2025',
     note='原图整数；9个来源地非完整世界总量；含不同煤种') for c,vs in countries.items() for y,v in zip([2023,2024,2025],vs)])
sources=json.loads((P/'sources.json').read_text(encoding='utf-8-sig'))
for source in sources:
    f=P/source['file']
    if f.exists():
        source.update(status='saved',bytes=f.stat().st_size,sha256=hashlib.sha256(f.read_bytes()).hexdigest())
        if source['id']=='EI_DATA_2025':
            source.pop('error',None)
            source['retrieval_method']='Chrome official archive download; urllib initially returned 403'
(P/'sources.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding='utf-8')
checks.append(dict(table='coal_origins_2023_2025_CCTD.csv',rows=27,method='transcribed from verified source PNG'))
(P/'validation.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False,indent=2))
