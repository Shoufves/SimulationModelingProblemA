"""Extract publisher-provided source workbooks without editing the originals."""
import csv
import hashlib
import json
import statistics
from collections import defaultdict
from pathlib import Path
import openpyxl

BASE = Path(__file__).resolve().parent
SRC = BASE/"extracted_source_data"/"Source_Data_files"
OUT = BASE/"data"
OUT.mkdir(exist_ok=True)

def write(name,header,rows):
    with (OUT/name).open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)

def main():
    w=openpyxl.load_workbook(SRC/"Source_Data_Fig. 4.xlsx",read_only=True,data_only=True)
    records=[]
    grouped=defaultdict(list)
    for source_row,row in enumerate(w.active.iter_rows(min_row=3,values_only=True),3):
        label,cost,income,year=row[:4]
        if not isinstance(year,(int,float)):
            continue
        residence,group=label.split(", ",1)
        records.append((source_row,residence,group,cost,income,year))
        grouped[(year,residence,group)].append((cost,income))
    write("household_energy_source_fig4.csv",
          ["source_excel_row","residence","author_income_group","energy_cost_RMB","household_income_RMB","year"],records)
    summary=[]
    for (year,residence,group),values in sorted(grouped.items()):
        costs=[v[0] for v in values]
        incomes=[v[1] for v in values]
        burdens=[100*c/i for c,i in values if i>0]
        summary.append((year,residence,group,len(values),statistics.mean(costs),
                        statistics.median(costs),statistics.median(incomes),
                        statistics.median(burdens),100*sum(b>10 for b in burdens)/len(burdens)))
    write("household_energy_by_author_group_unweighted.csv",
          ["year","residence","author_income_group","sample_records","mean_energy_cost_RMB",
           "median_energy_cost_RMB","median_household_income_RMB",
           "median_energy_burden_pct","share_burden_gt10_pct"],summary)
    w=openpyxl.load_workbook(SRC/"Source_Data_Fig. 5.xlsx",read_only=True,data_only=True)
    burdens=[]
    region=None
    for source_row,row in enumerate(w.active.iter_rows(min_row=4,values_only=True),4):
        if row[0] is not None:
            region=row[0]
        if row[1] is None:
            continue
        burdens.append((source_row,region,*row[1:6]))
    write("household_energy_burden_fig5.csv",
          ["source_excel_row","author_region_code","author_income_group","median_burden_2013_pct",
           "energy_poverty_share_2013_pct","median_burden_2017_pct","energy_poverty_share_2017_pct"],burdens)
    write("coal_employment_census_2023.csv",
          ["geography","statistical_period","industry","enterprise_legal_entities","employment_10000_persons","source_id","notes"],
          [("山西","2023年末","煤炭开采和洗选业",3000,93.27,"SX_CENSUS_TABLE32","人工核读原图表3-2"),
           ("内蒙古","2023年末","煤炭开采和洗选业",1077,19.9,"NM_CENSUS_2023","正文表3-2；原文2025-03-18，转载2025-05-21")])
    # Reclassify challenge responses; preserve them as evidence, never label them originals.
    manifest=json.loads((BASE/"sources_household_employment.json").read_text(encoding="utf-8"))
    for entry in manifest:
        p=BASE.parents[1]/entry["file"]
        if entry["id"]=="ENERGY_INEQUALITY_SOURCEDATA":
            entry.update(status="saved",download_method="Chrome公开下载；复制原下载文件",
                         bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest())
        elif entry["id"] in ("ENERGY_INEQUALITY_PAPER","ENERGY_INEQUALITY_SUPDATA"):
            entry.update(status="blocked_response",limitations=entry["limitations"]+"；程序收到3038字节拦截HTML，非论文/ZIP；Chrome页面已人工核实")
    (BASE/"sources_household_employment.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    derived=[]
    for p in sorted(list(OUT.glob("*.csv"))+list(SRC.glob("*.xlsx"))):
        derived.append(dict(file=str(p.relative_to(BASE)),bytes=p.stat().st_size,
                            sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
    (BASE/"derived_files_manifest.json").write_text(json.dumps(derived,ensure_ascii=False,indent=2),encoding="utf-8")
    print("Fig4 records",len(records),"summary groups",len(summary),"Fig5 rows",len(burdens))
    for row in summary:
        if row[0]==2017:
            print(row)

if __name__=="__main__":
    main()
