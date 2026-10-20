"""Audit saved research artifacts and build an offline reading guide."""
import csv, hashlib, json, shutil, logging
from pathlib import Path
from collections import Counter
from pypdf import PdfReader
from PIL import Image
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
logging.getLogger('pypdf').setLevel(logging.ERROR)
sources=json.loads((ROOT/'sources.json').read_text(encoding='utf-8'))
inventory=json.loads((ROOT/'数据文件目录.json').read_text(encoding='utf-8'))

def write(name,text):
    (ROOT/name).write_text(text.strip()+'\n',encoding='utf-8')

def link(path,label=None):
    return f'[{label or path}](<{(ROOT/path).as_posix()}>)'

# Keep user-provided competition materials with a separate local provenance record.
provided=ROOT/'raw/provided';provided.mkdir(exist_ok=True,parents=True)
original=ROOT.parent/'2026年第二届全国大学生仿真建模应用挑战赛赛题'/'A题'
provided_records=[]
for p in sorted(original.iterdir()):
    if p.suffix.lower() not in ('.pdf','.xlsx'): continue
    dest=provided/p.name
    if not dest.exists(): shutil.copy2(p,dest)
    item=dict(file=str(dest.relative_to(ROOT)).replace('\\','/'),original_file=str(p),sha256=hashlib.sha256(dest.read_bytes()).hexdigest(),bytes=dest.stat().st_size)
    if p.suffix=='.xlsx':
        wb=load_workbook(dest,read_only=True,data_only=True)
        item['sheets']=[dict(name=ws.title,nonempty_rows=sum(1 for row in ws.iter_rows(values_only=True) if any(v is not None for v in row)),max_columns=ws.max_column) for ws in wb]
        wb.close()
    provided_records.append(item)
write('附件检查.json',json.dumps(provided_records,ensure_ascii=False,indent=2))

checks=[]
for s in sources:
    if s['status']!='saved': continue
    p=ROOT/s['file']; row=dict(id=s['id'],file=s['file'],status='PASS',details='')
    try:
        b=p.read_bytes()
        assert len(b)==s['bytes'],'size mismatch'
        assert hashlib.sha256(b).hexdigest()==s['sha256'],'hash mismatch'
        if p.suffix.lower()=='.pdf':
            doc=PdfReader(p);row['details']=f'{len(doc.pages)} pages'
        elif p.suffix.lower() in ('.jpg','.jpeg','.png'):
            with Image.open(p) as im: im.verify()
        elif p.suffix.lower()=='.xlsx':
            wb=load_workbook(p,read_only=True,data_only=True);row['details']=' | '.join(wb.sheetnames);wb.close()
        elif p.suffix.lower()=='.xls':
            assert b.startswith(bytes.fromhex('D0CF11E0A1B11AE1')),'not XLS OLE signature'
            row['details']='XLS signature only; sheets not parsed'
    except Exception as e: row.update(status='FAIL',details=str(e))
    checks.append(row)
with (ROOT/'文件完整性检查.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,['id','file','status','details']);w.writeheader();w.writerows(checks)

# Text search copies, preserve PDF page numbers. Original PDFs remain authoritative.
for sid in ['CHINA_BTR1_2024','IRENA_COST_2024_FULL','IEA_CCUS_HYDROGEN','IEA_NEA_GENERATION_COST_2020']:
    s=next(s for s in sources if s['id']==sid); p=ROOT/s['file']
    target=ROOT/'extracted/reports'/f'{p.stem}.txt'
    if not target.exists():
        doc=PdfReader(p)
        target.write_text('\n\n'.join(f'--- PDF page {i+1} ---\n'+(pg.extract_text() or '') for i,pg in enumerate(doc.pages)),encoding='utf-8')

known={s['id'] for s in sources}
missing_data=[d['file'] for d in inventory if not (ROOT/d['file']).is_file()]
unmatched=[dict(file=d['file'],source=d['source']) for d in inventory if any(sid.strip() not in known for sid in d['source'].split(';'))]
stats=dict(date='2026-10-16',source_records=len(sources),saved=sum(s['status']=='saved' for s in sources),failed=sum(s['status']!='saved' for s in sources),raw_bytes=sum(s.get('bytes',0) for s in sources if s['status']=='saved'),derived_catalog_entries=len(inventory),integrity_failures=[c for c in checks if c['status']!='PASS'],missing_data=missing_data,source_id_review=unmatched,provided_files=len(provided_records))
write('核验摘要.json',json.dumps(stats,ensure_ascii=False,indent=2))

coverage=[
('问题一：能源产量、消费、占比','煤、石油、天然气及一次电力和其他','2005—2024年国家统计局能源生产/消费CSV；2016—2025年鉴能源平衡原表；2025年公报','核心总量已结构化；分品种平衡原图待提取','data/nbs/consumption_2005_2024.csv'),
('问题一：分技术装机','火电、水电、核电、风电、太阳能、生物质','2005—2024年官方装机；2016—2025年IRENA 15类装机；国家能源局2025年电力/可再生能源统计','主要装机序列已结构化；火电不等于煤电，生物质定义需统一','data/irena/china_capacity_2016_2025.csv'),
('问题一：地域分布','省级生产、用电、资源及电站位置','统计年鉴分地区表、赛题四附件；CEADs 1997—2015年省级能源清单','原件已下载；省级能源更新版2016—2022尚未取得','raw/provided'),
('问题一：进出口与安全','煤、油、气进出口和对外依存','历年年鉴主要商品进出口；2025年公报进口表；中国油气产业发展报告2026','原表已下载，2025进口表已提取；贸易伙伴结构未形成完整库','data/nbs'),
('问题一：价格','国际煤、油、气基准及国内价格、电价、碳价','世界银行年度/月度价格；国内2025年12月下旬流通价格；电力市场、碳市场报告及电价政策','国际序列已提取但需版本复核；国内连续年度价格仍待整理','data/prices'),
('问题二：需求与增长','GDP、人口、行业需求、能耗强度','年鉴GDP/人口/行业能源表；2005—2024年实际GDP增速CSV','GDP增速已结构化；分行业弹性需要估计','data/nbs/real_gdp_growth_2005_2024.csv'),
('问题二：历史校准与验证','全国/省级能源与碳排','官方能源序列、OWID中国能源与CO2；CarbonMonitor省级分部门日排放及年汇总','已提取，需统一边界；CarbonMonitor 2025仅到5月31日','data/emissions/carbonmonitor_province_sector_annual_2019_2025.csv'),
('问题二：排放与碳中和','燃料排放因子、电网因子、碳汇、CCUS','IPCC固定源燃烧；2023年中国电力因子；中国BTR1；IEA氢能与CCUS报告','电力因子已结构化；燃料与碳汇参数待按模型边界选择','data/emissions/china_grid_factors_2023.csv'),
('问题二：投资与技术演进','建设成本、运维、寿命、学习曲线、转化效率','IRENA成本报告；IEA/OECD发电成本；储能、氢能、可靠性报告；ABM论文','报告原件已下载，尚未建立经校准的完整参数表','raw/reports'),
('问题二/三：2030/2060基准与情景','政策目标、碳价、电价、技术进步、能源安全冲击','碳达峰方案、2035 NDC、十五五能源规划、容量电价、新能源电价、双碳白皮书','政策原文已下载；未来数值须明确设为情景假设','raw/policies'),
('问题三：极端事件与韧性','气候灾害、供电可靠性、进口冲击、储能','中国气候公报2025、电力可靠性2024、储能及油气报告','可支持年度/区域情景；逐小时负荷、出力、联络线数据尚未收集','raw/reports/climate_2025.pdf'),
]
write('01_题目需求与资料覆盖.md','# A题需求与资料覆盖\n\n整理日期：2026-10-16。下表区分“已下载原件”和“已整理为数据”；覆盖主题不代表每个指标、每个年份已齐全。\n\n|题目需求|具体内容|已取得资料|状态与缺口|本地入口|\n|---|---|---|---|---|\n'+'\n'.join('|'+ '|'.join([*r[:4],link(r[4],'打开')])+'|' for r in coverage))

write('02_数据口径与缺口.md', '''# 数据口径与缺口

建议先使用2015—2024年作为共同的十年历史窗口；2005—2024年用于较长周期校准。2025年更新数据单独核验后再加入，不把未满年的数据当全年。报告出版年不等于数据统计年。

## 必须保留的统计边界

- 全国能源主线优先采用国家统计局，电力技术细分参考国家能源局。OWID是二级汇编，保留codebook中的上游来源，与官方数据交叉验证，不直接拼接不同折算口径。
- 能源生产、终端消费、发电量、装机容量是不同变量。火电包含多类燃料；太阳能可能包含光热；氢是二次能源，储能是能量转移，不能重复计作一次能源产量。
- “清洁能源”与“非化石能源”不等价，天然气仍是化石能源。国家统计局“一次电力及其他能源”分类须保留原名。
- 国家统计局与国际数据库的一次电力折标方法可能不同；标准煤、实物量、热值与电量转换需记录系数及年份。
- IRENA可再生总装机不包含纯抽水蓄能；其生物能源/废弃物口径与国内生物质装机不同，不能不经解释相互替换。
- 中国CO2、能源CO2、化石燃料与水泥CO2、全部温室气体CO2e是不同边界。电力间接排放不能再叠加到已统计的发电侧排放。生物质也不能在所有边界中直接设为零排放。
- 森林蓄积量是存量，不是年度碳汇。碳中和需说明土地利用碳汇、CCUS、负排放与残余排放的计算边界。
- 单位GDP强度应统一不变价基期；年鉴能耗强度表变换基期后不能简单串接。名义价格与实际价格、美元与人民币、国际基准与国内到厂价必须分开。
- 原始年鉴档案含台湾地区能源表（例如C28-11），须独立标记地域，不能混入大陆全国序列。
- 历史年鉴可能存在版本修订，应优先选一致版本；缺失值保留为空，不用零代替。

## 已发现并保留的问题

1. 官方装机表2011年总量为106253万千瓦，分项加总106243万千瓦，相差10万千瓦；已对照原图，保留原值并记录差异。
2. CarbonMonitor原始标题覆盖2019—2025，实际日期为2019-01-01至2025-05-31。2025年155个省份×部门组合均标记为不完整年，禁止直接与完整年比较。
3. 世界银行价格工作簿含额外比较/差异工作表；提取实际价格页后，260组年度均价与12个月均值偏差均小于1%，但这不足以消除文件版本异常，数据标记为preliminary，正式论文引用前需复核版本。
4. OWID中国能源表2025年仅部分字段有值；中国CO2表实际最新为2024年。各字段可用年份见quality/owid_field_coverage.csv。
5. CEADs取得1997—2015年作者公开旧版XLS；只验证文件签名，未解析工作表。新版网页可访问，但2016—2022年文件下载接口尚未取得可用下载结果；不能宣称是完整最新省级清单。
6. 国家统计局历史图表保留了原图，多数尚未转为统一长表；报告文字提取用于检索，数值仍应回到PDF表格核对。

## 仍需补齐或由建模范围决定的内容

|项目|当前状态|后续处理|
|---|---|---|
|2016—2022年CEADs省级完整能源/碳清单|当前公开页面的下载操作未返回文件|继续定位作者镜像或使用省统计年鉴；省级部门排放可参考已取得的CarbonMonitor，但其方法和边界不同|
|煤油气国内连续年度价格、全国碳市场历史日价|取得部分报告统计、国际连续序列与国内时点价|需补取连续交易/统计序列；不得把时点价冒充年度均价|
|按贸易伙伴拆分的燃料进口、战略库存|已有总量与报告，完整细分未取得|若研究进口集中度需补贸易伙伴数据；未公开库存不能伪造|
|各技术CAPEX/OPEX、寿命、建设周期、效率|取得成本/技术报告|按中国适用性、价格年、技术边界提取为参数表并设敏感性区间|
|企业、居民、政府Agent决策参数|已有ABM方法文献，尚未估计|从历史数据校准行为弹性、投资规则，不能把主观设定称为观测值|
|逐小时负荷、风光出力、区域输电容量|尚未采集完整序列|年度系统演化模型可先采用利用小时与可靠性约束；若做小时调度再补采|

## 下载失败与替代

来源目录保留7条失败记录，便于追踪：国家统计局公报原地址失败但另一个官方地址及政府镜像已取得；IRENA成本网页受限但摘要与2024成本全文已取得；EI两个2026入口未取得，使用OWID历史汇编作补充且不冒充EI原件；海关页面失败，以年鉴及公报进口表补充；CEADs旧页面与figshare旧HTML失败，但作者公开API及1997—2015原件已取得。失败来源没有被计入成功下载数量。
''')

parameter_rows=[
('需求、GDP、人口、行业弹性',['NBS_2025_C09-02','NBS_2025_C03-04'],'年鉴能源消费与GDP指数；先估计历史需求弹性，人口/行业细分回到原表'),
('技术初始装机与新增/退役',['IRENA_CAPACITY_2026','NEA_POWER_2025','NEA_RENEWABLE_2025'],'IRENA容量报告中国行；官方电力统计；退役不能直接由净增装机推得'),
('可再生投资/运维与学习曲线',['IRENA_COST_2024_FULL','IRENA_COST_2025_SUMMARY'],'各技术总安装成本、LCOE、容量因子章节；全球均值与中国数值分开'),
('核电/常规发电经济参数',['IEA_NEA_GENERATION_COST_2020'],'核电、煤气、水电等技术章节；2020历史基准须更新价格与融资条件'),
('储能与氢能',['NEA_STORAGE_2025','NEA_HYDROGEN_2025','HYDROGEN_PLAN_2035'],'装机、时长、技术路线和应用场景；能量/功率、效率、利用率分开'),
('燃料和碳排放',['IPCC_COMBUSTION','CHINA_GRID_EF_2023'],'IPCC默认热值/排放因子及不确定性；中国电力因子用于相应核算边界'),
('CCUS、碳汇与残余排放',['IEA_CCUS_HYDROGEN','CHINA_BTR1_2024'],'CCUS路径及能耗；国家清单和土地利用/林业章节；存量与流量分开'),
('燃料价格与供给安全',['WB_PRICE_ANNUAL','NEA_OILGAS_2026'],'国际基准价格与油气供需；汇率、运输和税费另行处理'),
('碳市场、电价、容量收益',['MEE_CARBON_MARKET_2025','POWER_MARKET_2025','RENEWABLE_PRICE_136','CAPACITY_PRICE_2026'],'市场机制及实施时间；2023煤电容量政策须结合2026更新使用'),
('政策约束与情景目标',['POLICY_PEAK_2030','NDC_2035','ENERGY_PLAN_15_PDF'],'作为情景输入或外部比较标准；不能先强制达标再将结果称为预测'),
('气候冲击与供电可靠性',['CMA_CLIMATE_2025','RELIABILITY_2024'],'历史灾害、设备可靠性；需把报告统计映射成冲击概率/损失区间'),
('投资主体行为与交互',['ABM_INVESTMENT_PAPER','ABM_CHINA_RIGHTS_PAPER'],'参考方法与机制设计；论文参数不能未经校准直接移植中国全能源系统'),
]
byid={s['id']:s for s in sources}
write('03_Mesa建模参数资料索引.md','# Mesa建模参数资料索引\n\n暂按Python + Mesa组织。下表是参数取数入口，尚不是已校准参数表。每个正式参数应保留：名称、Agent类型、值、单位、价格年/统计年、地区、来源页码、区间、处理方法。\n\n|参数组|本地来源|使用方法|\n|---|---|---|\n'+'\n'.join('|'+name+'|'+'；'.join(link(byid[i]['file'],byid[i]['title']) for i in ids if i in byid)+'|'+note+'|' for name,ids,note in parameter_rows)+'\n\n实施顺序建议：先跑通全国年度能源平衡，再加入区域差异、企业投资与政策反馈；基准情景和至少三个对照情景使用同一历史校准参数集，仅调整明确列出的外生假设。成本、排放、供需平衡和安全指标应同时输出。')

selected=[s for s in sources if s['status']=='saved' and not s['id'].startswith(('NBS_ARCHIVE','NBS_FRAME','NBS_MENU','FIGSHARE_FILE'))]
write('04_核心资料索引.md','# 核心资料索引\n\n原始网址、获取日期、SHA256和失败记录见资料来源目录.csv / sources.json。以下链接均指向本地副本。\n\n'+'\n'.join(f"- **{s['agency']}**：{link(s['file'],s['title'])}。覆盖：{s['coverage']}。来源ID：`{s['id']}`。" for s in selected))

write('00_阅读入口.md',f'''# A题资料库阅读入口

整理日期：2026-10-16。服务于“双碳背景下多种能源融合发展系统仿真与优化研究”，暂按Python + Mesa建模。

已保存 **{stats['saved']}份公开来源文件**（含目录网页、统计表图片和报告，不等于{stats['saved']}套独立数据），原始下载约 **{stats['raw_bytes']/1024/1024:.1f} MiB**；另归档赛题PDF与4份附件。来源目录保留 **{stats['failed']}条失败下载记录**及错误信息。已整理 **{len(inventory)}项数据/字典/检查文件**，原始资料、处理结果和质量说明分别保存。

题目要求的主要资料类别已有来源，但目前仍不是所有指标都齐全、统一口径的最终建模数据集。尤其省级更新清单、国内连续价格和贸易伙伴结构仍有缺口。

## 从这里开始

1. {link('01_题目需求与资料覆盖.md','题目需求与资料覆盖')}：按题目三问逐项检查已有资料及缺口。
2. {link('02_数据口径与缺口.md','数据口径与缺口')}：建模前必读，避免跨口径拼接和重复计量。
3. {link('数据文件目录.csv','已整理数据目录')}：21项文件包含核心CSV、字段说明和检查表，具体以目录内容为准。
4. {link('03_Mesa建模参数资料索引.md','Mesa参数取数入口')}：成本、效率、投资、排放、政策和风险分别去哪找。
5. {link('04_核心资料索引.md','核心原始资料索引')} / {link('资料来源目录.csv','完整来源目录')}：查看本地原件、来源机构、年份和下载地址。

## 可优先使用的数据

|内容|时间覆盖|文件|
|---|---|---|
|官方能源消费及结构|2005—2024|{link('data/nbs/consumption_2005_2024.csv','消费CSV')}|
|官方能源生产及结构|2005—2024|{link('data/nbs/production_2005_2024.csv','生产CSV')}|
|官方分技术发电装机|2005—2024|{link('data/nbs/capacity_2005_2024.csv','装机CSV')}|
|官方实际GDP增速|2005—2024|{link('data/nbs/real_gdp_growth_2005_2024.csv','GDP CSV')}|
|IRENA中国15类技术装机|2016—2025|{link('data/irena/china_capacity_2016_2025.csv','分技术CSV')}|
|省级分部门排放年汇总|2019—2024完整；2025至5月31日|{link('data/emissions/carbonmonitor_province_sector_annual_2019_2025.csv','排放CSV')}|
|中国电力排放因子|2023|{link('data/emissions/china_grid_factors_2023.csv','因子CSV')}|

国际能源价格、OWID全国能源和CO2补充数据见数据目录。价格文件带版本复核标记；用于论文前应先读质量说明。2025年国家统计局公报、能源局统计已保存，但并未把所有2025字段强行并入历史一致口径序列。

## 文件组织与复核

- `raw/`：下载原件；`raw/provided/`：比赛提供材料的副本。
- `data/`：已提取CSV及质量检查；`extracted/`：报告检索文本，数值以原PDF为准。
- `sources.json`：完整来源元数据；`jobs_*.json`：下载清单；`scripts/`：采集和处理脚本。
- {link('文件完整性检查.csv','文件完整性检查')}：已保存文件逐一核对大小与SHA256；PDF检查可读取页数，图片验证格式，XLSX打开检查。旧XLS只检查签名。
- {link('核验摘要.json','核验摘要')} / {link('附件检查.json','比赛附件检查')}：机器可读审计结果。

## 复现说明

在资料库目录运行：`python scripts/collect.py jobs_initial.json`（其他jobs文件同理）重新下载；脚本优先使用本地缓存。整理数据依次运行 `python scripts/prepare_data.py`、`python scripts/enrich_data.py`，最后运行 `python scripts/finalize_library.py` 更新检查和入口。需要Python、openpyxl、pypdf、pdfplumber、Pillow。不要只运行prepare后就把目录当作完整结果；enrich会补入分技术装机、省级排放、电力因子和GDP数据。官方图片中手工录入的值也保存在处理脚本内，并非自动OCR。

引用时按具体原始表/报告注明机构、年份和页码，不把本资料库视为原始数据发布者。报告是参考资料，未来目标和情景假设不属于历史观测。
''')
print(json.dumps(stats,ensure_ascii=False,indent=2))
