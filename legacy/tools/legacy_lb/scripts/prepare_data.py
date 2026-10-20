"""Reproducible extraction. Originals remain untouched; blanks remain missing."""
import csv,json,re,hashlib,statistics
from pathlib import Path
from collections import defaultdict,Counter
from html.parser import HTMLParser
import openpyxl
from pypdf import PdfReader
from collect import ROOT,decode

OUT=ROOT/'data';OUT.mkdir(exist_ok=True)
inventory=[];checks=[]
def write(name,rows,fields=None,source='',note=''):
    rows=list(rows);fields=fields or (list(rows[0]) if rows else [])
    p=OUT/name;p.parent.mkdir(parents=True,exist_ok=True)
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fields,extrasaction='ignore');w.writeheader();w.writerows(rows)
    inventory.append(dict(file=str(p.relative_to(ROOT)).replace('\\','/'),rows=len(rows),columns=len(fields),source=source,note=note))

# Visually transcribed from official 2025 yearbook tables, retained in source units.
consumption='''2005,261369,72.4,17.8,2.4,7.4
2006,284647,72.4,17.5,2.7,7.4
2007,311442,72.5,17.0,3.0,7.5
2008,320611,71.5,16.7,3.4,8.4
2009,336126,71.6,16.4,3.5,8.5
2010,360648,69.2,17.4,4.0,9.4
2011,387043,70.2,16.8,4.6,8.4
2012,402138,68.5,17.0,4.8,9.7
2013,416913,67.4,17.1,5.3,10.2
2014,428334,65.8,17.3,5.6,11.3
2015,434113,63.8,18.4,5.8,12.0
2016,441492,62.2,18.7,6.1,13.0
2017,455827,60.6,18.9,6.9,13.6
2018,471925,59.0,18.9,7.6,14.5
2019,485584,57.5,19.0,8.0,15.4
2020,496514,56.7,18.9,8.4,16.0
2021,525027,55.7,18.8,8.8,16.7
2022,541274,55.6,18.4,8.4,17.6
2023,571327,54.8,18.8,8.5,17.9
2024,596000,53.2,18.2,8.8,19.8'''
production='''2005,229037,77.4,11.3,2.9,8.4
2006,244763,77.5,10.8,3.2,8.5
2007,264173,77.8,10.1,3.5,8.6
2008,277419,76.8,9.8,3.9,9.5
2009,286092,76.8,9.4,4.0,9.8
2010,312125,76.2,9.3,4.1,10.4
2011,340178,77.8,8.5,4.1,9.6
2012,351041,76.2,8.5,4.1,11.2
2013,358784,75.4,8.4,4.4,11.8
2014,362212,73.5,8.3,4.7,13.5
2015,362193,72.2,8.5,4.8,14.5
2016,345954,69.8,8.3,5.2,16.7
2017,358867,69.6,7.6,5.4,17.4
2018,378859,69.2,7.2,5.4,18.2
2019,392606,68.1,7.0,5.7,19.2
2020,402671,67.1,6.9,6.1,19.9
2021,422789,66.3,6.7,6.2,20.8
2022,459507,66.9,6.4,6.0,20.8
2023,475979,66.1,6.3,6.1,21.5
2024,498000,63.9,6.1,6.2,23.8'''
capacity='''2005,51718,39138,11739,696,106,,
2006,62370,48382,13029,696,207,,
2007,71822,55607,14823,908,420,,
2008,79273,60286,17260,908,839,,
2009,87410,65108,19629,908,1760,3,3
2010,96641,70967,21606,1082,2958,26,3
2011,106253,76834,23298,1257,4623,212,19
2012,114676,81968,24947,1257,6142,341,20
2013,125768,87009,28044,1466,7652,1589,8
2014,137887,93232,30486,2008,9657,2486,19
2015,152527,100554,31954,2717,13075,4218,9
2016,165051,106094,33207,3364,14747,7631,7
2017,177708,110495,34359,3582,16325,12942,7
2018,189997,114408,35259,4466,18427,17433,5
2019,200985,118957,35804,4874,20915,20429,5
2020,220504,124960,37028,4989,28165,25356,5
2021,237689,129739,39094,5326,32871,30654,5
2022,256317,133527,41396,5557,36564,39268,5
2023,292224,139099,42237,5691,44144,61048,5
2024,334922,144462,43552,6083,52078,88742,5'''

for key,text,source in [('consumption',consumption,'NBS_2025_C09-02'),('production',production,'NBS_2025_C09-01')]:
    fields=['year',key+'_10ktce','coal_share_pct','oil_share_pct','gas_share_pct','primary_electricity_and_other_share_pct']
    rows=[]
    for vals in csv.reader(text.splitlines()):
        row=dict(zip(fields,[int(vals[0]),int(vals[1])]+list(map(float,vals[2:]))));row['source_id']=source;rows.append(row)
        checks.append(dict(check=key+'_shares_sum',year=row['year'],value=round(sum(row[k] for k in fields[2:]),3),status='PASS' if abs(sum(row[k] for k in fields[2:])-100)<=0.11 else 'REVIEW'))
    write('nbs/'+key+'_2005_2024.csv',rows,source=source,note='Visually transcribed; total 10,000 tonnes standard coal; shares percent; original primary electricity and other category preserved')
rows=[];fields=['year','total_10kW','thermal_10kW','hydro_10kW','nuclear_10kW','wind_10kW','solar_10kW','other_10kW']
for vals in csv.reader(capacity.splitlines()):
    row=dict(zip(fields,[int(v) if v else '' for v in vals]));row['source_id']='NBS_2025_C09-15';rows.append(row)
    if all(row[k]!='' for k in fields[2:]):
        diff=sum(row[k] for k in fields[2:])-row['total_10kW'];checks.append(dict(check='capacity_components_minus_total_10kW',year=row['year'],value=diff,status='PASS_ROUNDING' if abs(diff)<=3 else 'REVIEW'))
write('nbs/capacity_2005_2024.csv',rows,source='NBS_2025_C09-15',note='Thermal includes coal/gas/oil/biomass; solar not identical to PV; blanks are missing, not zero')

# OWID: separate compilations, without conversion or filling missing values.
coverage=[]
for dataset in ['energy','co2']:
    with (ROOT/f'raw/owid/owid-{dataset}-data.csv').open(encoding='utf-8-sig') as f:
        reader=csv.DictReader(f);fields=reader.fieldnames;rows=[r for r in reader if r['country']=='China' and 2000<=int(r['year'])<=2025]
    write(f'owid/china_{dataset}_2000_2025_available.csv',rows,fields,source='OWID_'+('ENERGY' if dataset=='energy' else 'CO2'),note='Original field names/units. 2025 may be partial or absent. See coverage and codebook. Not national standard-coal accounting.')
    codebook=list(csv.DictReader((ROOT/f'raw/owid/owid-{dataset}-codebook.csv').open(encoding='utf-8-sig')))
    write(f'owid/{dataset}_codebook.csv',codebook,source='OWID_'+('ENERGY' if dataset=='energy' else 'CO2')+'_CODEBOOK',note='Includes upstream sources, descriptions and units')
    for field in fields:
        valid=[int(r['year']) for r in rows if r[field] not in ('',None)]
        target=[int(r['year']) for r in rows if 2015<=int(r['year'])<=2024 and r[field] not in ('',None)]
        coverage.append(dict(dataset=dataset,field=field,first_year=min(valid) if valid else '',last_year=max(valid) if valid else '',available_years=len(valid),available_years_2015_2024=len(target),missing_years_2015_2024=';'.join(str(y) for y in range(2015,2025) if y not in target)))
write('quality/owid_field_coverage.csv',coverage,source='OWID_ENERGY;OWID_CO2')

# World Bank data: record exact source cell, source units and cache values only.
price_rows={}
for mode,sheet,header_row,unit_row in [('Annual','Annual Prices (Nominal)',7,8),('Monthly','Monthly Prices',5,6)]:
    wb=openpyxl.load_workbook(ROOT/f'raw/worldbank/CMO_Historical_{mode}.xlsx',read_only=True,data_only=True);ws=wb[sheet]
    head=[ws.cell(header_row,c).value for c in range(1,12)];units=[ws.cell(unit_row,c).value for c in range(1,12)]
    rows=[]
    for ri,vals in enumerate(ws.iter_rows(min_row=unit_row+1,max_col=11,values_only=True),unit_row+1):
        period=vals[0]
        valid=(isinstance(period,(int,float)) and 2000<=period<=2025) if mode=='Annual' else bool(re.match(r'^20(?:0\d|1\d|2[0-6])M\d\d$',str(period)))
        if not valid:continue
        for j in range(1,11):
            if not isinstance(vals[j],(int,float)):continue
            rows.append(dict(period=int(period) if mode=='Annual' else period,commodity=head[j],value=vals[j],unit=units[j],source_id='WB_PRICE_'+mode.upper(),sheet=sheet,cell=f'{openpyxl.utils.get_column_letter(j+1)}{ri}',quality_status='preliminary_source_workbook_has_comparison_sheets'))
    price_rows[mode]=rows
    write(f'prices/worldbank_energy_{mode.lower()}_2000_latest.csv',rows,source='WB_PRICE_'+mode.upper(),note='Nominal USD prices. Downloaded workbook includes comparison/mismatch sheets; extracted source cells preserved and annual/monthly checks provided. Do not treat this as China domestic price.')
annual={(r['period'],r['commodity']):r['value'] for r in price_rows['Annual']};monthly=defaultdict(list)
for r in price_rows['Monthly']:
    k=(int(r['period'][:4]),str(r['commodity']).replace(' **',''));monthly[k].append(r['value'])
comparison=[]
for k,value in annual.items():
    vals=monthly.get(k,[])
    if len(vals)==12:
        mean=statistics.mean(vals);rel=abs(value-mean)/max(abs(value),1e-9)
        comparison.append(dict(year=k[0],commodity=k[1],annual=value,monthly_mean=mean,relative_difference=rel,status='PASS_1PCT' if rel<0.01 else 'REVIEW'))
write('quality/worldbank_annual_monthly_check.csv',comparison,source='WB_PRICE_ANNUAL;WB_PRICE_MONTHLY',note='Simple monthly means are a consistency diagnostic, not proof of correct annual weighting')

# Preserve bulletin tables as source-table extraction, not silently harmonized model input.
class Tables(HTMLParser):
    def __init__(self):super().__init__();self.tables=[];self.depth=0;self.rows=None;self.row=None;self.cell=None
    def handle_starttag(self,t,a):
        if t=='table':
            self.depth+=1
            if self.depth==1:self.rows=[]
        elif self.depth==1 and t=='tr':self.row=[]
        elif self.depth==1 and t in ('td','th'):self.cell=[]
    def handle_data(self,d):
        if self.cell is not None:self.cell.append(d)
    def handle_endtag(self,t):
        if self.depth==1 and t in ('td','th') and self.cell is not None:
            if self.row is not None:self.row.append(re.sub(r'\s+',' ',''.join(self.cell)).strip())
            self.cell=None
        elif self.depth==1 and t=='tr' and self.row is not None:
            self.rows.append(self.row);self.row=None
        elif t=='table':
            if self.depth==1 and self.rows:self.tables.append(self.rows)
            self.depth=max(0,self.depth-1)
parser=Tables();parser.feed(decode((ROOT/'raw/nbs/2025_bulletin_official.html').read_bytes()))
for i,table in enumerate(parser.tables,1):
    if not any(any(k in c for k in ['原煤','原油','天然气','发电','煤及','能源']) for row in table for c in row):continue
    fields=['col_'+str(j+1) for j in range(max(map(len,table)))];rows=[dict(zip(fields,row)) for row in table]
    write(f'nbs/bulletin2025_source_table_{i:02d}.csv',rows,fields,source='NBS_2025_BULLETIN_ALTERNATE',note='Source table extraction; original merged headers/footnotes must be checked in archived HTML')

# Reports searchable by PDF page number, not replacing originals.
for p in list((ROOT/'raw/reports').glob('*.pdf'))+list((ROOT/'raw/policies').glob('*.pdf')):
    dest=ROOT/'extracted/reports'/(p.stem+'.txt');dest.parent.mkdir(parents=True,exist_ok=True)
    if dest.exists():continue
    try:
        reader=PdfReader(p);dest.write_text('\n'.join(f'\n=== PDF page {i+1} ===\n'+(pg.extract_text() or '') for i,pg in enumerate(reader.pages)),encoding='utf-8')
    except Exception as e:checks.append(dict(check='pdf_extract',file=p.name,status='REVIEW',error=str(e)))

# CarbonMonitor full original retained; filtered 2019-2025 rows and column metadata.
p=ROOT/'raw/carbonmonitor/CM_China_v1.csv'
if p.exists():
    with p.open(encoding='utf-8-sig') as f:
        reader=csv.DictReader(f);fields=reader.fieldnames;rows=list(reader)
    inventory.append(dict(file=str(p.relative_to(ROOT)).replace('\\','/'),rows=len(rows),columns=len(fields),source='FIGSHARE_FILE_55875995',note='Original near-real-time estimates; preserve source time/province/sector/unit. Do not equate with annual inventory.'))
    (ROOT/'extracted/carbonmonitor_schema.json').write_text(json.dumps({'columns':fields,'rows':len(rows),'first_rows':rows[:3]},ensure_ascii=False,indent=2),encoding='utf-8')

write('quality/nbs_transcription_checks.csv',checks,fields=sorted(set(k for r in checks for k in r)),source='NBS_2025_C09-01;NBS_2025_C09-02;NBS_2025_C09-15')
(ROOT/'数据文件目录.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
with (ROOT/'数据文件目录.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,['file','rows','columns','source','note']);w.writeheader();w.writerows(inventory)
print(json.dumps({'data_files':len(inventory),'total_rows':sum(x['rows'] for x in inventory),'review_checks':[x for x in checks if x.get('status')=='REVIEW'],'price_checks':dict(Counter(x['status'] for x in comparison))},ensure_ascii=False))
