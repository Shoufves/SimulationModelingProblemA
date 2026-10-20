import json,re,hashlib,csv,logging
from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
P=Path(__file__).resolve().parent
logging.getLogger('pypdf').setLevel(logging.ERROR)
(P/'extracted').mkdir(exist_ok=True)
sources=json.loads((P/'sources.json').read_text(encoding='utf-8'))
checks=[]
for s in sources:
    if s['status']!='saved':continue
    p=P/s['file'];b=p.read_bytes();c={'id':s['id'],'file':s['file'],'sha256_match':hashlib.sha256(b).hexdigest()==s['sha256']}
    if p.suffix=='.pdf':
        d=PdfReader(p);texts=[pg.extract_text() or '' for pg in d.pages];c['pages']=len(texts)
        (P/'extracted'/f'{p.stem}.txt').write_text('\n\n'.join(f'--- PDF page {i+1} ---\n'+t for i,t in enumerate(texts)),encoding='utf-8')
        print(s['id'],len(texts),'pages',texts[0][:220].replace('\n',' '))
    elif p.suffix=='.xlsx':
        w=load_workbook(p,read_only=True,data_only=True);c['sheets']=[]
        for ws in w:
            rows=list(ws.values);c['sheets'].append({'name':ws.title,'rows':len(rows),'columns':ws.max_column})
            out=P/'extracted'/f'{p.stem}_{re.sub(r"[^a-zA-Z0-9_-]","_",ws.title)}.csv'
            with out.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(rows)
            print(s['id'],ws.title,rows[:7])
        w.close()
    elif p.suffix=='.html':
        text=p.with_suffix('.txt').read_text(encoding='utf-8')
        c['text_characters']=len(text)
        print(s['id'],'text length',len(text))
    checks.append(c)
(P/'audit.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
