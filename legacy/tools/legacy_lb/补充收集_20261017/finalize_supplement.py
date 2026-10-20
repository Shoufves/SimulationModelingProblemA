import json,csv,hashlib,re
from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
P=Path(__file__).resolve().parent; LIB=P.parent
def link(rel,label=None):return f'[{label or rel}](<{(P/rel).as_posix()}>)'
def write(name,text):(P/name).write_text(text.strip()+'\n',encoding='utf-8')
sources=[]
for s in json.loads((P/'international/sources.json').read_text(encoding='utf-8')):
    s['file']='international/'+s['file'];s['group']='international';sources.append(s)
for s in json.loads((P/'ai_social/sources_ai_social.json').read_text(encoding='utf-8')):
    s['file']=str((LIB/s['file']).relative_to(P)).replace('\\','/');s['coverage']=s.get('data_years','');s['notes']=s.get('use','')+'；'+s.get('limitations','');s['group']='ai_social';sources.append(s)
old=json.loads((LIB/'sources.json').read_text(encoding='utf-8'));old_hash={s.get('sha256'):s['id'] for s in old if s.get('sha256')}
audit=[]
for s in sources:
    if s['status']!='saved':continue
    f=P/s['file'];b=f.read_bytes();check={'id':s['id'],'file':s['file'],'bytes_ok':len(b)==s['bytes'],'sha256_ok':hashlib.sha256(b).hexdigest()==s['sha256']}
    if f.suffix=='.pdf':check['pdf_pages']=len(PdfReader(f).pages)
    if f.suffix=='.xlsx':
        w=load_workbook(f,read_only=True,data_only=True);check['sheets']=w.sheetnames;w.close()
    if s['sha256'] in old_hash:check['same_bytes_as_existing_id']=old_hash[s['sha256']]
    audit.append(check)
write('sources_supplement.json',json.dumps(sources,ensure_ascii=False,indent=2))
with (P/'资料来源目录.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,['id','group','title','agency','publication_date','coverage','url','file','status','bytes','sha256','notes','error'],extrasaction='ignore');w.writeheader();w.writerows(sources)
stats={'date':'2026-10-17','records':len(sources),'saved':len(audit),'failed':sum(s['status']!='saved' for s in sources),'bytes':sum(s.get('bytes',0) for s in sources if s['status']=='saved'),'checks':audit,'integrity_failures':[a for a in audit if not a['bytes_ok'] or not a['sha256_ok']],'duplicate_existing':[a for a in audit if a.get('same_bytes_as_existing_id')]}
write('核验摘要.json',json.dumps(stats,ensure_ascii=False,indent=2))
write('01_情景与资料映射.md',f'''# 情景与资料映射

S0—S5作为本轮确定的实验框架推进；社会支持配对实验和S6协调应对保留，具体措施仍需由结果识别。此处仅提供取数依据，尚未赋予全部参数数值。

|情景|本轮重点资料|模型变量/规则|仍需估计或校准|
|---|---|---|---|
|S0 基准|2026碳配额方案、数据中心行动计划、需求响应办法、2025—2027储能方案；EIA2026年10月展望；NBS居民收入|行业覆盖、配额分配、设施效率约束、居民组别；初始国际市场条件|已宣布目标如何影响主体决策；2030后政策延续规则；经济与人口路径|
|S1 政策强化|碳配额方案PDF及通知；CBAM正式阶段官方网页与中文解释|改变行业配额规则、基准值或其他政策工具，观察企业响应|政策强度范围、财政预算与执行程度；不能把内生碳价同时当外生指定值|
|S2 国际风险|EIA六个通道XLSX、能源安全季度表、美国LNG资料、欧盟能源退出安排|进口暴露、通道可用量、替代供应、运输成本和恢复路径|中国分来源贸易矩阵、可动用库存、实际替代能力、冲击概率与持续期|
|S3 技术进步|本轮储能行动方案；复用既有IRENA成本全文、IEA氢能/CCUS和发电成本报告|技术成本、效率、建设时间、融资与部署约束|中国适用参数、价格年调整、技术学习路径；政策目标不当观测值|
|S4 AI快速发展|IEA2025/2026、四部委行动计划、信通院指标定义、国内报道交叉参考|算力容量、利用率、PUE、选址、任务可调比例、电力需求|大陆基期口径、AI专属份额、任务弹性；IEA China含香港不可直接套用|
|S5 联合冲击|复用S2及S4相同参数与随机样本|计算联合效应与单独效应的差异|不额外随意改变GDP、天气等，避免把混杂效应当交互效应|
|社会支持配对/S6|NBS分组收入、世界银行中国/山西报告、煤矿关闭报告、ILO案例、需求响应办法|居民负担、岗位转移、培训与收入支持、政府预算、灵活负荷|家庭能源账单、分省最新就业、转岗成功率与时滞；案例参数不能当全国均值|

## 首先应修正的模型设想

1. 2026配额方案正文强调以碳排放强度控制实施免费分配，发电行业使用基准法，钢铁/水泥/铝冶炼2026年度使用强度系数法（PDF页1—3）。基准不能无说明地写成固定绝对排放总量下的统一拍卖市场。
2. 算力的上架率、IT负载率、PUE、设备算力不是同一个指标；不能把EFLOPS唯一换算为电量。
3. 居民消费“居住”项不是能源账单，不能直接作为能源负担分子。
4. 企业退出涉及的岗位不等于净失业人数；退休、转岗、新增岗位和空间错配需分开。
5. 国际通道数据存在同一货物经过多个海峡的情况，不能把通道流量相加当独立进口来源。

详细来源可从{link('资料来源目录.csv','来源目录')}按ID检索。
''')
write('02_口径问题与未解决缺口.md',f'''# 口径问题与未解决缺口

## 必须先解决的口径

- **研究范围是中国大陆。** IEA Energy and AI相关报告中的China包含香港，不能直接初始化主模型。子目录已保存区域定义核验图；不能随意减去香港全部用电以替代同口径数据中心用电。
- **算力设施、数据中心与AI数据中心不是同一统计集合。** 国内报道与IEA估计不可直接拼接。国内报道只称“据相关机构统计”，已降为交叉参考，不称官方调查数据。
- **实际、估计、预测与政策目标分开。** EIA 2026年10月展望发布于10月6日，输入截至10月1日，预测至2027；并不反映10月17日所有新事件。2026整年值仍非全年实绩。
- **中国库存数据的地位。** EIA能源安全表包含其对中国战略石油库存的估计，不能称中国官方公布，也不能直接等同模型可动用应急库存。库存能力、实际库存和释放速率分别建模。
- **政策目标不当成已实现能力。** 需求响应办法的2025年3%—5%最大负荷目标、储能2027年1.8亿千瓦以上目标均应标记政策目标。不能把目标写成实测或无条件强制模型达标。
- **电力直接排放与用电间接排放不得重复计算。** 配额方案中铝冶炼包含CF4/C2F6等，需区分CO2与CO2e指标。
- **进出口边际表不是双边贸易矩阵。** 同一通道按来源与按目的地的两张表，不足以确定某来源国到中国的双边流量。

## 已发现的原始文件问题

1. EIA能源安全2026年8月表2的马六甲2025Q4总油流量为24.9，表3同项为24.0，单位均为百万桶/日。已保留两值并在quality_notes.json记录，暂不自动修正或混用。
2. EIA通道专题HTML下载件只有页面框架；正文已通过Chrome读取，正式定量使用同页六个原始XLSX附件。该HTML不宣称完整离线报告。
3. 商务部网站CBAM中文手册URL含202511，但当前PDF封面为2026年7月，机构为商务部外贸发展事务局；以正文版次记录。它是解释材料，法律规则以欧盟正式文件为准。
4. EIA STEO工作簿原样保存，混合历史、估计、预测；部分未来单元格为零占位，不能一律解释成实际零值。提取的整表CSV仅便于检索，不是已清洗统一参数表。
5. 全球通道流量不能直接替换中国进口流量；1H25是半年平均，不能当全年观测。

## 仍缺的关键参数

|参数|现有支撑|缺口与处理|
|---|---|---|
|大陆AI专属用电与区域分布|国际估计、官方政策、国内交叉报道|定义不一致，需同口径基期；暂列校准/敏感性参数|
|任务转移比例与算力选址偏好|技术报告提供机制|缺全国代表性观测，不预设所有AI任务可延后|
|收入组能源账单|NBS收入与消费分类|缺收入组×城乡×能源品种交叉账单，居住消费不能代替|
|煤炭地区就业流量|历史案例与研究|缺近期分省就业、退休、招聘、转岗率，不把旧值直接初始化|
|中国进口来源、储备与替代供给|既有年鉴/油气报告、本轮通道和国际供应数据|尚非完整双边流量矩阵；战略库存仅有外部估计|
|国际冲击概率与持续时间|历史流量变化和短期市场报告|没有单一权威概率；应作为清晰定义的压力测试参数|
|建设、技术成本与价格传导|既有成本报告及本轮政策|中国本土参数及价格年、传导系数仍需校准|

## 无法取得的公开页面

本轮4个入口未下载成功：IEA中东专题、IEA2026年10月石油市场页、欧盟理事会原公告、UNCTAD2025海运报告。保留403失败记录；未越过浏览器验证。欧盟规则已有欧委会替代官方页，近期市场背景已有EIA2026年10月完整展望及表格，通道数据已有EIA原始附件。未取得的文件不计作成功。
''')
write('03_原件与数据索引.md','# 原件与数据索引\n\n'+ '\n'.join(f"- `{s['id']}` **{s['status']}**："+(link(s['file'],s['title']) if s['status']=='saved' else s['title'])+f"；{s['agency']}；发布：{s.get('publication_date','')}；覆盖：{s.get('coverage','')}。{s.get('notes','')}" for s in sources))
write('00_阅读入口.md',f'''# 第二轮资料补充：情景与主体参数

保存和核验日期：2026-10-17。本轮由主智能体与一个用户指定的GPT-6.1-sol子智能体分工完成。

已保存 **{len(audit)}份来源原件/网页**，共约 **{stats['bytes']/1024/1024:.1f} MiB**；另保留 **{stats['failed']}条未成功下载的来源记录**。原有资料库未被覆盖，本轮清单独立存放。此数量包含政策通知、解释页与附件，不等于独立数据库数量。

## 阅读顺序

1. {link('01_情景与资料映射.md','情景与模型参数对应')}：S0—S6各用什么资料，哪些变量仍需校准。
2. {link('02_口径问题与未解决缺口.md','口径问题与缺口')}：尤其注意IEA中国含香港、居民居住支出不是能源账单、外部库存估计不是中国官方值。
3. {link('03_原件与数据索引.md','原件索引')} / {link('资料来源目录.csv','完整来源目录')}：本地文件、网址、机构、年份及限制。
4. {link('ai_social/00_阅读入口.md','AI与社会专题说明')}：算力、居民收入及就业转型的详细取数入口。

## 本轮重点补充

- 国际风险：EIA海峡/航道原始Excel、能源安全季度数据、2026年10月短期能源展望及全表、美国LNG、欧盟能源进口与CBAM规则。
- 政策与协调措施：中国2026年碳配额方案、2025—2027储能建设行动方案、需求响应办法、数据中心绿色低碳行动计划。
- AI：IEA2025与2026报告、信通院统计指标定义、国内算力报道交叉参考。
- 社会：2024/2025居民收入消费与五等份收入，世界银行中国/山西及煤矿关闭报告、ILO中国企业转型案例。
- 技术成本：继续复用第一轮已取得的IRENA、IEA氢能/CCUS及发电成本报告，本轮不重复下载。

## 已整理的数据

|文件|内容|使用限制|
|---|---|---|
|{link('international/data/route_crude_origin_destination_2020_1H2025.csv','航道来源/去向长表')}|636行，六条通道、2020—2025上半年，保留原单元格|不是双边贸易矩阵，通道不可简单相加|
|{link('international/data/china_destination_route_flows_2020_1H2025.csv','以China为目的地的航道流量')}|36行，来自同一原始图表|保留原地理标签，需与大陆海关口径核对|
|{link('international/data/eia_security_quarterly_2025_2026.csv','能源安全季度表')}|279行，库存与通道估计；保留原始缺失符号|外部估计；存在一项跨表冲突，见质量说明|

整表转出的CSV在international/extracted，保持原表结构，仅用于检索。核心原始数值不因口径冲突被擅自改写。国际冲击、AI增长和社会支持的最终参数表尚未完成校准。

## 核验与复现

{link('核验摘要.json','核验摘要')}记录每个成功文件的大小、SHA256、PDF页数或Excel工作表。失败下载不计入成功数量。数据转换脚本保留在international目录；AI/社会来源采集脚本保留在对应子目录。

本轮资料支持确定情景参数的来源与取值依据，不代表所有参数均有公开观测值，更不代表模型或情景实验已完成。
''')
print(json.dumps({k:v for k,v in stats.items() if k!='checks'},ensure_ascii=False))
