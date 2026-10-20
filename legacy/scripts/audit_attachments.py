"""Read the competition attachments without changing the source workbooks."""

from collections import Counter
import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]


def audit_workbook(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for sheet in workbook:
        records = [
            (index, list(row))
            for index, row in enumerate(sheet.iter_rows(values_only=True), 1)
            if index > 1 and any(value is not None for value in row)
        ]
        names = Counter(str(row[0]).strip() for _, row in records if row[0] is not None)
        capacities = [
            row[1] for _, row in records if isinstance(row[1], (int, float))
        ]
        summary = {
            "sheet": sheet.title,
            "formatted_max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "header": list(next(sheet.iter_rows(values_only=True))),
            "records": len(records),
            "unique_names": len(names),
            "capacity_sum_raw": sum(capacities),
            "capacity_min_max_raw": [min(capacities), max(capacities)],
            "duplicate_names": {
                name: [
                    {"row": index, "values": row}
                    for index, row in records
                    if str(row[0]).strip() == name
                ]
                for name, count in names.items()
                if count > 1
            },
            "blank_cells_by_column": {
                str(column + 1): sum(
                    row[column] is None or str(row[column]).strip() == ""
                    for _, row in records
                )
                for column in range(sheet.max_column)
            },
            "storage_named_records": [
                {"row": index, "values": row}
                for index, row in records
                if "抽水" in str(row[0])
            ],
            "unnamed_column_values": [
                {"cell": sheet.cell(index, column + 1).coordinate, "value": row[column]}
                for index, row in records
                for column, header in enumerate(next(sheet.iter_rows(values_only=True)))
                if header is None and row[column] is not None
            ],
            "first_records": records[:2],
            "last_records": records[-2:],
        }
        if path.name == "附件1.xlsx":
            summary["kw_conversion_mismatches"] = [
                {"row": index, "values": row[:4]}
                for index, row in records
                if isinstance(row[1], (int, float))
                and isinstance(row[3], (int, float))
                and abs(row[3] - 1000 * row[1]) > 0.01
            ]
        sheets.append(summary)
    formula_workbook = openpyxl.load_workbook(path, data_only=False)
    formulas = [
        {"sheet": sheet.title, "cell": cell.coordinate, "formula": cell.value}
        for sheet in formula_workbook
        for row in sheet
        for cell in row
        if cell.data_type == "f"
    ]
    comments = [
        {"sheet": sheet.title, "cell": cell.coordinate, "text": cell.comment.text}
        for sheet in formula_workbook
        for row in sheet
        for cell in row
        if cell.comment
    ]
    return {"file": path.name, "sheets": sheets, "formulas": formulas, "comments": comments}


def main():
    result = [
        audit_workbook(path)
        for path in sorted(ROOT.glob("附件*.xlsx"))
    ]
    output = ROOT / "outputs" / "attachments" / "yx_附件数据核查.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    for book in result:
        for sheet in book["sheets"]:
            print(json.dumps({
                "file": book["file"],
                "sheet": sheet["sheet"],
                "records": sheet["records"],
                "unique_names": sheet["unique_names"],
                "capacity_sum_raw": sheet["capacity_sum_raw"],
                "duplicate_names": list(sheet["duplicate_names"]),
            }, ensure_ascii=False))
    print(f"Full audit: {output}")


if __name__ == "__main__":
    main()
