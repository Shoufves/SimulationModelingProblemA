"""Build the user manual from its Markdown source using bundled python-docx."""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/13_workbench_user_manual.md"
OUTPUT = ROOT / "outputs/manuals/多能源仿真实验台操作手册.docx"


def text(paragraph, value):
    """Preserve literal names without Markdown syntax in the Word document."""
    paragraph.add_run(value.replace("`", ""))


def table(doc, lines):
    records = [[cell.strip() for cell in line.strip("|").split("|")] for line in lines]
    records = [row for row in records if not all(re.fullmatch(r"[:\- ]+", cell) for cell in row)]
    count = len(records[0])
    widths = {2: (2.3, 4.6), 3: (2.9, 2, 2), 4: (2.1, .6, 1.2, 3.0)}[count]
    result = doc.add_table(rows=0, cols=count)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for col, width in zip(result.columns, widths):
        col.width = Inches(width)
    props = result._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        for name, value in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
            edge.set(qn("w:" + name), value)
        borders.append(edge)
    props.append(borders)
    for index, record in enumerate(records):
        row = result.add_row()
        # Keep individual rows intact while allowing the table itself to flow.
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for col, (cell, value, width) in enumerate(zip(row.cells, record, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, size in (("top", "80"), ("bottom", "80"), ("left", "95"), ("right", "95")):
                item = OxmlElement("w:" + side)
                item.set(qn("w:w"), size)
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            tcpr.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "DDE8EF" if index == 0 else ("F5F7F9" if index % 2 == 0 else "FFFFFF"))
            tcpr.append(shading)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(14)
            if count == 4 and col in (1, 2) or count == 3 and col in (1, 2):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            run.font.size = Pt(10.5)
            run.bold = index == 0
    spacer = doc.add_paragraph()
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.6)
    sec.left_margin = sec.right_margin = Inches(.8)
    sec.footer_distance = Inches(.25)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3",
                 "Caption", "List Number", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = Pt(16)
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        style.element.get_or_add_pPr().append(snap)
        for border in list(style.element.xpath("./w:pPr/w:pBdr")):
            border.getparent().remove(border)
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            style.element.rPr.rFonts.attrib.pop(qn("w:" + attr), None)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = Pt(16)
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["Title"].paragraph_format.line_spacing = Pt(31)
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    for name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.line_spacing = Pt(size + 6)
        doc.styles[name].paragraph_format.space_before = Pt(9)
        doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = False
    doc.styles["Caption"].font.bold = False
    doc.styles["Caption"].paragraph_format.line_spacing = Pt(13)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    next_page = False
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line == "<!-- pagebreak -->":
            next_page = True
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            p = doc.add_paragraph(line[3:], "Heading 1")
            p.paragraph_format.page_break_before = next_page
            next_page = False
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            # Crop complete dashboard sections to keep labels readable in print.
            source_image = SOURCE.parent / match[2]
            crop_height = {"agents-desktop.png": 1100, "operations-desktop.png": 1100,
                           "evidence-desktop.png": 1010, "export-desktop.png": 610}[source_image.name]
            crop_dir = ROOT / "tmp/workbench-manual-images"
            crop_dir.mkdir(parents=True, exist_ok=True)
            cropped = crop_dir / source_image.name
            with Image.open(source_image) as image:
                image.crop((0, 50, image.width, min(crop_height, image.height))).save(cropped)
            p.paragraph_format.line_spacing = 1
            picture = p.add_run().add_picture(str(cropped), width=Inches(6.9))
            picture._inline.docPr.set("descr", match[1])
        elif line.startswith("|"):
            block = [line]
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index].strip())
                index += 1
            table(doc, block)
        elif line.startswith("```"):
            while index < len(lines) and not lines[index].strip().startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(lines[index])
                r.font.name = "Consolas"
                r.font.size = Pt(10)
                index += 1
            index += 1
        elif re.match(r"\d+\. ", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.18)
            p.paragraph_format.first_line_indent = Inches(-.18)
            text(p, line)
        else:
            p = doc.add_paragraph(style="Caption" if re.match(r"图\d ", line) else "Normal")
            text(p, line)
    doc.core_properties.title = "多能源仿真实验台操作手册"
    doc.core_properties.subject = "本地实验台操作与结果留存"
    doc.core_properties.author = "A题项目组"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
